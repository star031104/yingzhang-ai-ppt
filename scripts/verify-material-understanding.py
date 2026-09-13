"""Reproducible upload-to-export check using clearly labeled synthetic source files."""
import io
import json
import os
import sys
from pathlib import Path
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "test-results" / "material-understanding"
OUTPUT.mkdir(parents=True, exist_ok=True)
os.environ["SLIDEFORGE_DATABASE_URL"] = f"sqlite:///{OUTPUT.as_posix()}/verification.db"
os.environ["SLIDEFORGE_ARTIFACT_ROOT"] = str(OUTPUT / "artifacts")
os.environ["SLIDEFORGE_PUBLIC_TEST_MODE"] = "false"
sys.path.insert(0, str(ROOT / "apps" / "api"))

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app


def save_bytes(document):
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def fixtures():
    doc = Document()
    doc.add_heading("知识问答试点复盘", 0)
    doc.add_paragraph("合成验证材料，用于测试文档理解和演示排版，不是真实业务数据。")
    sections = [
        ("业务背景", ["知识分散在制度文件与操作手册中，检索后仍需人工判断版本。", "业务人员需要得到附带原始出处的回答，便于快速复核。"]),
        ("核心问题", ["现有检索能够找到相似内容，却难以区分过期规定与当前口径。", "对于缺乏证据的问题，系统必须明确说明材料不足。"]),
        ("系统架构", ["资料索引保存文件版本、章节与页码，回答能够定位到原始段落。", "检索层只提供与问题相关的片段，复核人员能够查看完整上下文。", "回答附带引用，原文未提供的信息保持未知。"]),
        ("实施方法", ["首先整理现行文件，标记发布日期与适用业务。", "然后将典型问题与原始证据配对，建立复核样本。", "最后由业务人员逐条确认答案和引用，再逐步扩大试点。"]),
    ]
    for title, paragraphs in sections:
        doc.add_heading(title, 1)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
    doc.add_heading("试点验证结果", 1)
    doc.add_paragraph("表 1 同一批内部问答样本的复核结果")
    table = doc.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ["方案", "准确率（%）", "召回率（%）"]):
        cell.text = value
    for values in [["关键词检索", "71.2", "68.4"], ["语义检索", "84.5", "82.1"], ["检索与人工复核", "89.6", "87.3"]]:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    doc.add_paragraph("结果表明检索与人工复核的组合改善了回答质量。")
    doc.add_heading("结论与边界", 1)
    for text in ["建议优先在内部知识问答场景推广，并保留人工复核。", "当前结果仅限内部样本，不能推断其他行业或公开服务中的效果。", "扩大试点前，应补充异常问题与过期文件的评估。"]:
        doc.add_paragraph(text)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "服务指标"
    sheet.append(["方案", "平均耗时（秒）", "复核覆盖率"])
    for row in [["关键词检索", 32, .70], ["语义检索", 24, .82], ["检索与人工复核", 18, .96]]:
        sheet.append(row)
        sheet.cell(sheet.max_row, 3).number_format = "0%"
    return [("试点复盘（合成材料）.docx", save_bytes(doc)), ("服务指标（合成数据）.xlsx", save_bytes(workbook))]


def main():
    with TestClient(app) as client:
        project = client.post("/api/v1/projects", json={"name": "材料理解与排版验证"}).json()
        base = f"/api/v1/projects/{project['id']}"
        for name, data in fixtures():
            (OUTPUT / name).write_bytes(data)
            uploaded = client.post(base + "/sources", files={"file": (name, data)})
            assert uploaded.status_code == 201, uploaded.text
            assert uploaded.json()["readingStatus"] == "read"
        response = client.post(base + "/outline", json={"title": "让每个回答都有据可查", "preset": "business", "slide_count": 9,
                              "instructions": "根据复盘报告与服务指标形成管理层汇报，重点讲清方案、结果与适用边界。", "skill_ids": ["data-consulting"]})
        assert response.status_code == 200, response.text
        plan = response.json()
        (OUTPUT / "plan.json").write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
        text = "\n".join(point for slide in plan["slides"] for point in slide["content"]["bullets"])
        for value in ["89.6%", "87.3%", "18秒", "不能推断其他行业"]:
            assert value in text, f"Source information missing: {value}"
        generated = client.post(base + "/generate")
        assert generated.status_code == 200, generated.text
        workspace = client.get(base + "/workspace").json()
        checked = client.post(base + "/validate")
        assert checked.status_code == 200, checked.text
        quality = checked.json()
        (OUTPUT / "quality-report.json").write_text(json.dumps(quality, ensure_ascii=False, indent=2), encoding="utf-8")
        assert quality["passed"] and quality["blockingErrors"] == 0, quality
        render_root = Path(generated.json()["html"]).parent
        selected = []
        for slide in workspace["slides"]:
            current = json.loads((render_root / "slides" / str(slide["position"]) / "current.json").read_text(encoding="utf-8"))
            score = current["scoreDetail"]
            for field in ["overflow", "textOverflow", "missingContent", "lowContrast", "missingAssets"]:
                assert score.get(field, 0) == 0, (slide["position"], field, score)
            selected.append({"position": slide["position"], "variant": current["variant"], "score": current["score"],
                             "image": str(render_root / "slides" / str(slide["position"]) / f"{current['variant']}.png")})
        for extension in ["pptx", "html", "pdf"]:
            exported = client.get(base + f"/export/{extension}")
            assert exported.status_code == 200, exported.text[:300]
            (OUTPUT / f"材料理解验证.{extension}").write_bytes(exported.content)
        with ZipFile(OUTPUT / "材料理解验证.pptx") as archive:
            parts = [name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
            package_text = "\n".join(archive.read(name).decode() for name in parts)
            assert len(parts) == 9
            assert "不能推断其他行业" in package_text
            charts = len([name for name in archive.namelist() if name.startswith("ppt/charts/chart") and name.endswith(".xml")])
            tables = package_text.count("<a:tbl>")
            assert charts >= 1 and tables >= 1, (charts, tables)
        report = {"projectId": project["id"], "slides": selected, "nativeCharts": charts, "nativeTables": tables,
                  "qualityPassed": quality["passed"], "qualitySuggestions": quality.get("issues", []),
                  "sourceFiles": [source["name"] for source in workspace["sources"]],
                  "verification": "Actual API uploads, rules-only generation, Chromium previews, editable PPTX objects and PDF export. Synthetic data; desktop Office rendering not covered."}
        (OUTPUT / "verification.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
