import io
import re
from collections import Counter, defaultdict
from itertools import pairwise
from statistics import median

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

FUNCTION_VARIANTS = {
    "cover": ["hero", "editorial-cover", "minimal-cover"],
    "agenda": ["numbered-list", "section-map", "agenda-cards"],
    "section": ["chapter-divider", "section-statement", "minimal-section"],
    "data": ["chart-focus", "table-highlight", "metric-wall"],
    "comparison": ["comparison-bars", "two-column", "scorecard"],
    "process": ["workflow", "three-stage", "process"],
    "image-story": ["figure-wide", "figure-analysis", "image-story"],
    "conclusion": ["takeaways", "summary-grid", "next-steps"],
    "content": ["split", "statement", "cards"],
}


def _title_text(slide) -> str:
    for shape in slide.shapes:
        if not getattr(shape, "is_placeholder", False):
            continue
        if shape.placeholder_format.type in {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}:
            return str(getattr(shape, "text", "")).strip()
    return next(
        (
            str(shape.text).strip()
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and str(shape.text).strip()
        ),
        "",
    )


def _slide_function(index: int, title: str, counts: Counter, text_length: int) -> str:
    normalized = re.sub(r"\s+", "", title).lower()
    if index == 1 and text_length < 180:
        return "cover"
    if re.search(r"目录|议程|agenda|contents?", normalized, re.IGNORECASE):
        return "agenda"
    if re.search(r"总结|结论|建议|下一步|conclusion|summary|nextsteps?", normalized, re.IGNORECASE):
        return "conclusion"
    if re.search(r"对比|对照|优势|vs\.?|comparison", normalized, re.IGNORECASE):
        return "comparison"
    if counts["chart"] or counts["table"]:
        return "data"
    if counts["connector"] >= 2 or re.search(
        r"流程|路径|方法|架构|process|workflow|architecture",
        normalized,
        re.IGNORECASE,
    ):
        return "process"
    if counts["picture"] and counts["picture_area"] > 0.18:
        return "image-story"
    if text_length < 90 and counts["text"] <= 3:
        return "section"
    return "content"


def _silhouette(slide, width: int, height: int, counts: Counter) -> dict:
    content_shapes = [
        shape
        for shape in slide.shapes
        if not getattr(shape, "is_placeholder", False)
        or getattr(shape.placeholder_format, "type", None)
        not in {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
    ]
    left = sum((shape.left + shape.width / 2) < width / 2 for shape in content_shapes)
    right = len(content_shapes) - left
    title_shape = next(
        (
            shape
            for shape in slide.shapes
            if getattr(shape, "is_placeholder", False)
            and shape.placeholder_format.type in {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
        ),
        None,
    )
    title_band = "none"
    if title_shape:
        center = (title_shape.top + title_shape.height / 2) / height
        title_band = "top" if center < 0.28 else "center" if center < 0.68 else "bottom"
    balance = (
        "balanced"
        if abs(left - right) <= max(1, len(content_shapes) * 0.25)
        else "left-heavy"
        if left > right
        else "right-heavy"
    )
    visible_shapes = [
        shape for shape in slide.shapes if int(shape.width) > 0 and int(shape.height) > 0
    ]
    left_edges = [round(int(shape.left) / width, 3) for shape in visible_shapes]
    right_edges = [
        round((width - int(shape.left + shape.width)) / width, 3) for shape in visible_shapes
    ]
    top_edges = [round(int(shape.top) / height, 3) for shape in visible_shapes]
    bottom_edges = [
        round((height - int(shape.top + shape.height)) / height, 3) for shape in visible_shapes
    ]
    alignment_x = Counter(round(value / 0.025) * 0.025 for value in left_edges)
    return {
        "titleBand": title_band,
        "columnBalance": balance,
        "objectCount": len(slide.shapes),
        "textCount": counts["text"],
        "mediaAreaRatio": round(counts["picture_area"], 3),
        "hasChart": bool(counts["chart"]),
        "hasTable": bool(counts["table"]),
        "hasDiagram": counts["connector"] >= 2,
        "safeMargins": {
            "left": max(0.0, min(left_edges, default=0.04)),
            "right": max(0.0, min(right_edges, default=0.04)),
            "top": max(0.0, min(top_edges, default=0.04)),
            "bottom": max(0.0, min(bottom_edges, default=0.04)),
        },
        "alignmentColumns": [round(value, 3) for value, _count in alignment_x.most_common(5)],
    }


def _constraint_grammar(slides: list[dict]) -> dict:
    silhouettes = [slide["silhouette"] for slide in slides]
    safe_margins = [item["safeMargins"] for item in silhouettes]
    media_ratios = [float(item["mediaAreaRatio"]) for item in silhouettes]
    object_counts = [int(item["objectCount"]) for item in silhouettes]
    columns = Counter(value for item in silhouettes for value in item.get("alignmentColumns", []))
    return {
        "version": "template-constraint-grammar-v1",
        "safeMargins": {
            side: round(median([item[side] for item in safe_margins]), 3)
            for side in ("left", "right", "top", "bottom")
        },
        "titleBands": Counter(item["titleBand"] for item in silhouettes).most_common(),
        "columnBalance": Counter(item["columnBalance"] for item in silhouettes).most_common(),
        "mediaAreaRange": [
            round(min(media_ratios, default=0.0), 3),
            round(max(media_ratios, default=0.0), 3),
        ],
        "objectCountRange": [
            min(object_counts, default=0),
            max(object_counts, default=0),
        ],
        "alignmentGrid": [round(value, 3) for value, _count in columns.most_common(8)],
        "functionSilhouettes": {
            function: Counter(
                (
                    slide["silhouette"]["titleBand"],
                    slide["silhouette"]["columnBalance"],
                )
                for slide in slides
                if slide["function"] == function
            ).most_common()
            for function in sorted({slide["function"] for slide in slides})
        },
    }


def analyze_reference_pptx(data: bytes) -> dict:
    deck = Presentation(io.BytesIO(data))
    fonts, sizes, colors = Counter(), Counter(), Counter()
    layouts, functions = Counter(), Counter()
    layout_by_function: dict[str, Counter] = defaultdict(Counter)
    slides = []
    width, height = int(deck.slide_width), int(deck.slide_height)
    for index, slide in enumerate(deck.slides, 1):
        layout_name = slide.slide_layout.name
        layouts[layout_name] += 1
        counts: Counter = Counter()
        text_length = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                counts["picture"] += 1
                counts["picture_area"] += (shape.width * shape.height) / max(1, width * height)
            if getattr(shape, "has_chart", False):
                counts["chart"] += 1
            if getattr(shape, "has_table", False):
                counts["table"] += 1
            if shape.shape_type in {MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.GROUP}:
                counts["connector"] += 1
            if not getattr(shape, "has_text_frame", False):
                continue
            counts["text"] += 1
            text_length += len(str(shape.text or ""))
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.name:
                        fonts[run.font.name] += 1
                    if run.font.size:
                        sizes[round(run.font.size.pt)] += 1
                    if run.font.color.type and run.font.color.rgb:
                        colors[str(run.font.color.rgb)] += 1
        title = _title_text(slide)
        function = _slide_function(index, title, counts, text_length)
        functions[function] += 1
        layout_by_function[function][layout_name] += 1
        slides.append(
            {
                "position": index,
                "title": title[:160],
                "function": function,
                "layout": layout_name,
                "silhouette": _silhouette(slide, width, height, counts),
                "textLength": text_length,
            }
        )
    sequence = [slide["function"] for slide in slides]
    transitions = Counter(f"{left}→{right}" for left, right in pairwise(sequence))
    exemplars = {
        function: [slide["position"] for slide in slides if slide["function"] == function][:3]
        for function in functions
    }
    variant_by_role = {
        "cover": FUNCTION_VARIANTS["cover"],
        "agenda": FUNCTION_VARIANTS["agenda"],
        "section": FUNCTION_VARIANTS["section"],
        "data": FUNCTION_VARIANTS["data"],
        "comparison": FUNCTION_VARIANTS["comparison"],
        "method": FUNCTION_VARIANTS["process"],
        "architecture": ["layered-architecture", "pipeline", "hub-spoke"],
        "evidence": ["evidence-chain", "source-map", "claim-map"],
        "conclusion": FUNCTION_VARIANTS["conclusion"],
        "content": FUNCTION_VARIANTS["content"],
    }
    return {
        "kind": "visual-skill",
        "version": "reference-analysis-v3",
        "fontFamily": fonts.most_common(3),
        "fontScale": sizes.most_common(8),
        "palette": colors.most_common(8),
        "layoutPatterns": layouts.most_common(),
        "functionPatterns": functions.most_common(),
        "functionLayoutMap": {
            key: value.most_common() for key, value in layout_by_function.items()
        },
        "rhythm": {
            "sequence": sequence,
            "transitions": transitions.most_common(),
            "exemplars": exemplars,
        },
        "variantByRole": variant_by_role,
        "constraintGrammar": _constraint_grammar(slides),
        "slides": slides,
        "slideSize": {"width": width, "height": height},
    }
