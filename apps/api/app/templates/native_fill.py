import io
import zipfile
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER


TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE}
ROLE_LAYOUT_HINTS = {
    "cover": ("title", "cover", "封面", "标题"),
    "section": ("section", "chapter", "章节", "节标题"),
    "data": ("chart", "data", "图表", "数据"),
    "comparison": ("comparison", "two", "对比", "比较"),
    "architecture": ("diagram", "process", "架构", "流程"),
    "method": ("process", "content", "方法", "流程"),
    "questions": ("ending", "closing", "结束", "致谢"),
}


def _placeholder_types(layout) -> list:
    return [shape.placeholder_format.type for shape in layout.placeholders]


def _layout_score(layout, spec: dict) -> float:
    role = str(spec.get("role", "content"))
    name = str(layout.name or "").lower()
    types = _placeholder_types(layout)
    title_count = sum(item in TITLE_TYPES for item in types)
    body_count = sum(item in BODY_TYPES for item in types)
    picture_count = sum(item == PP_PLACEHOLDER.PICTURE for item in types)
    score = title_count * 8 + min(body_count, 2) * 5
    score += sum(12 for hint in ROLE_LAYOUT_HINTS.get(role, ()) if hint.lower() in name)
    has_image = any(
        item.get("type") in {"source-image", "generated-image"} and Path(item.get("path", "")).is_file()
        for item in spec.get("assetBindings", [])
    )
    if has_image:
        score += picture_count * 20
    if role in {"cover", "section", "questions"}:
        score += 12 if title_count and body_count <= 1 else -body_count * 3
    elif not body_count:
        score -= 35
    if role in {"data", "comparison", "architecture", "method"} and body_count:
        score += 8
    if "blank" in name or "空白" in name:
        score -= 30
    return score


def _select_layout(deck, spec: dict):
    return max(deck.slide_layouts, key=lambda layout: _layout_score(layout, spec))


def _fill_text_frame(shape, values: list[str]) -> None:
    frame = shape.text_frame
    frame.clear()
    for index, value in enumerate(values):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = str(value)
        paragraph.level = 0


def _bound_image(spec: dict) -> Path | None:
    for item in spec.get("assetBindings", []):
        candidate = Path(item.get("path", ""))
        if item.get("type") in {"source-image", "generated-image"} and candidate.is_file():
            return candidate
    return None


def _restore_fidelity_parts(template: bytes, output: Path) -> int:
    prefixes = ("ppt/theme/", "ppt/slideMasters/", "ppt/slideLayouts/")
    with zipfile.ZipFile(io.BytesIO(template)) as source_package:
        source_parts = {
            name: source_package.read(name)
            for name in source_package.namelist()
            if name.startswith(prefixes) and name.endswith(".xml")
        }
    with zipfile.ZipFile(output) as result_package:
        result_parts = {
            name: result_package.read(name) for name in result_package.namelist()
        }
    result_parts.update(source_parts)
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as result_package:
        for name, payload in result_parts.items():
            result_package.writestr(name, payload)
    return len(source_parts)


def fill_native_template(template: bytes, slides: list[dict], output: Path) -> dict:
    result = Presentation(io.BytesIO(template))
    while len(result.slides) > 0:
        slide_id = result.slides._sldIdLst[-1]
        result.part.drop_rel(slide_id.rId)
        del result.slides._sldIdLst[-1]
    filled = 0
    body_filled = 0
    image_filled = 0
    layout_usage: Counter = Counter()
    unmatched: list[int] = []
    for spec in slides:
        layout = _select_layout(result, spec)
        layout_usage[layout.name] += 1
        slide = result.slides.add_slide(layout)
        title_done = body_done = image_done = False
        image = _bound_image(spec)
        for shape in slide.placeholders:
            placeholder_type = shape.placeholder_format.type
            if placeholder_type in TITLE_TYPES and not title_done:
                shape.text = spec.get("content", {}).get("title", spec.get("message", ""))
                title_done = True
            elif placeholder_type == PP_PLACEHOLDER.PICTURE and image and not image_done:
                shape.insert_picture(str(image))
                image_done = True
            elif placeholder_type in BODY_TYPES and not body_done:
                bullets = list(spec.get("content", {}).get("bullets", []))
                if not bullets and spec.get("message"):
                    bullets = [spec["message"]]
                _fill_text_frame(shape, bullets)
                body_done = True
        filled += int(title_done)
        body_filled += int(body_done)
        image_filled += int(image_done)
        if not title_done or (spec.get("role") not in {"cover", "section", "questions"} and not body_done):
            unmatched.append(int(spec.get("position", len(unmatched) + 1)))
    output.parent.mkdir(parents=True, exist_ok=True)
    result.save(output)
    fidelity_parts = _restore_fidelity_parts(template, output)
    return {
        "slides": len(slides),
        "titlesFilled": filled,
        "bodiesFilled": body_filled,
        "imagesFilled": image_filled,
        "layoutUsage": dict(layout_usage),
        "unmatchedSlides": unmatched,
        "mode": "semantic-layout-routing-v2",
        "fidelityPartsRestored": fidelity_parts,
        "themeMasterRoundtrip": "byte-stable",
    }
