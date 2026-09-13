import csv
import hashlib
import io
import re
from collections import Counter
from pathlib import Path

import fitz
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from app.documents.structure import parse_docx_structure, table_record, table_text
from app.intelligence.content_selection import section_understanding

PARSER_VERSION = "5.5"


def _aligned_metric_tables(page, captions):
    """Recover simple borderless numeric tables, only with aligned complete rows."""
    from types import SimpleNamespace
    lines = []
    for block in page.get_text("dict").get("blocks", []):
        for line in block.get("lines", []):
            text = "".join(span.get("text", "") for span in line.get("spans", [])).strip()
            if text:
                lines.append((fitz.Rect(line["bbox"]), text))
    lines.sort(key=lambda item: (round(item[0].y0 / 3), item[0].x0))
    groups = []
    for rect, text in lines:
        if not groups or abs(groups[-1][0][0].y0 - rect.y0) > 3:
            groups.append([])
        groups[-1].append((rect, text))
    found = []
    numeric = re.compile(r"^[−+\-]?\d+(?:\.\d+)?%?$")
    for caption in captions:
        start_y = caption["bbox"][3]
        candidates = [sorted(group, key=lambda cell: cell[0].x0) for group in groups if group[0][0].y0 > start_y]
        if not candidates:
            continue
        header = candidates[0]
        if not 2 <= len(header) <= 6 or header[0][0].y0 - start_y > 45:
            continue
        if any(len(text) > 30 or numeric.fullmatch(text) for _, text in header):
            continue
        rows = [header]
        ambiguous = False
        for row in candidates[1:]:
            if row[0][0].y0 - rows[-1][0][0].y0 > 45:
                break
            if len(row) != len(header):
                ambiguous = len(row) > 1
                break
            if any(abs((rect.x0 + rect.x1) / 2 - (head.x0 + head.x1) / 2) > 20 for (rect, _), (head, _) in zip(row, header)):
                ambiguous = True
                break
            if not all(numeric.fullmatch(text) for _, text in row[1:]) or len(row[0][1]) > 40:
                ambiguous = True
                break
            rows.append(row)
        if ambiguous or len(rows) < 3:
            continue
        bbox = fitz.Rect(rows[0][0][0])
        for row in rows:
            for rect, _ in row:
                bbox |= rect
        values = [[text for _, text in row] for row in rows]
        found.append(SimpleNamespace(bbox=tuple(bbox), header=SimpleNamespace(external=False), extract=lambda values=values: values))
    return found

SEMANTIC_PATTERNS = [
    ("conclusion", re.compile(r"结论|总结|展望|局限|创新点|conclusion|future work|limitation", re.IGNORECASE)),
    ("data", re.compile(r"实验|结果|评估|性能|指标|数据集|消融|财务|现金流|销售预测|收入|资金|资本结构|experiment|result|evaluation|metric|ablation", re.IGNORECASE)),
    ("architecture", re.compile(r"架构|框架|系统设计|总体设计|技术路线|产品概况|产品结构|角色协同|architecture|framework|system design", re.IGNORECASE)),
    ("method", re.compile(r"方法|算法|流程|模型|知识库|检索|提示词|产品服务|功能模块|市场营销|推广|运营|method|algorithm|workflow|model|retrieval|prompt", re.IGNORECASE)),
    ("comparison", re.compile(r"相关工作|研究现状|国内外|对比|基线|竞品|竞争环境|行业分析|related work|state of the art|baseline|comparison", re.IGNORECASE)),
    ("problem", re.compile(r"问题|挑战|目标|任务|贡献|缺口|主要工作|研究内容|风险|problem|challenge|objective|contribution|gap", re.IGNORECASE)),
    ("background", re.compile(r"摘要|引言|背景|意义|绪论|项目介绍|商机|政策|abstract|introduction|background|motivation", re.IGNORECASE)),
]

FIGURE_CAPTION_RE = re.compile(
    r"^\s*((?:图|Figure)\s*\d+(?:\s*[-－.]\s*\d+)?(?:\s*[（(][a-zA-Z0-9]+[）)])?)\s*(.*)$",
    re.IGNORECASE,
)
TABLE_CAPTION_RE = re.compile(
    r"^\s*((?:表|Table)\s*\d+(?:\s*[-－.]\s*\d+)?)\s*(.*)$",
    re.IGNORECASE,
)


def _section(index: int, title: str, text: str, page: int | None = None) -> dict:
    return {"id": f"S{index:03d}", "title": title, "text": text.strip(), "page": page}


def _clean_caption(value: str) -> str:
    text = re.sub(r"\s+", " ", value or "").strip(" ：:")
    return text.replace("处理时间不同任务的平均", "不同任务的平均处理时间")


PAGE_MARKER_RE = re.compile(
    r"^(?:第\s*\d+\s*页|\d+\s*/\s*\d+|[-—–]?\s*\d+\s*[-—–]?)$",
    re.IGNORECASE,
)


def _is_page_marker(value: str) -> bool:
    return bool(PAGE_MARKER_RE.fullmatch(_clean_caption(value)))


def _is_contents_page(value: str) -> bool:
    """Identify real contents pages so their dot leaders never become evidence."""
    lines = [_clean_caption(line) for line in str(value or "").splitlines() if line.strip()]
    if not lines:
        return False
    leaders = sum(bool(re.search(r"[.．·…]{5,}\s*\d*\s*$", line)) for line in lines[1:])
    chapter_rows = sum(bool(re.match(r"^(?:第\s*[一二三四五六七八九十\d]+\s*章|[（(][一二三四五六七八九十\d]+[）)])", line)) for line in lines[1:])
    has_heading = bool(re.fullmatch(r"目\s*录|contents?", lines[0], re.IGNORECASE))
    return leaders >= 5 or (has_heading and (leaders >= 3 or chapter_rows >= 5))


def _pdf_visual_headings(page: fitz.Page) -> set[str]:
    """Recover headings from font hierarchy when PDF text has no structure tags."""
    rows: list[tuple[float, str]] = []
    for block in page.get_text("dict").get("blocks", []):
        for line in block.get("lines", []):
            spans = line.get("spans", [])
            text = _clean_caption("".join(str(span.get("text", "")) for span in spans))
            if not text or _is_page_marker(text):
                continue
            rows.append((max((float(span.get("size", 0)) for span in spans), default=0), text))
    if not rows:
        return set()
    common_size = Counter(round(size, 1) for size, text in rows if len(text) >= 6).most_common(1)
    body_size = common_size[0][0] if common_size else 0
    return {
        text for size, text in rows
        if size >= body_size + 1.2 and 2 <= len(text) <= 70
        and not FIGURE_CAPTION_RE.match(text) and not TABLE_CAPTION_RE.match(text)
    }


def _pdf_page_title(page: fitz.Page, page_number: int) -> str:
    candidates: list[tuple[float, float, str]] = []
    for block in page.get_text("dict").get("blocks", []):
        if "lines" not in block:
            continue
        for line in block["lines"]:
            spans = line.get("spans", [])
            text = _clean_caption("".join(str(span.get("text", "")) for span in spans))
            if (
                not 2 <= len(text) <= 90
                or _is_page_marker(text)
                or FIGURE_CAPTION_RE.match(text)
                or TABLE_CAPTION_RE.match(text)
            ):
                continue
            y = float(line.get("bbox", [0, 0, 0, 0])[1])
            if y < page.rect.height * 0.03 or y > page.rect.height * 0.88:
                continue
            size = max((float(span.get("size", 0)) for span in spans), default=0)
            heading_like = bool(re.match(r"^(?:第\s*[一二三四五六七八九十\d]+\s*[章节]|\d+(?:\.\d+){0,2}\s+|摘要|Abstract|结论|总结|参考文献|致谢)", text, re.IGNORECASE))
            candidates.append((1 if heading_like else 0, size - y / max(1, page.rect.height), text))
    if not candidates:
        return f"第 {page_number} 页"
    heading_like = [item for item in candidates if item[0] == 1]
    selected = max(heading_like or candidates, key=lambda item: (item[0], item[1]))
    return selected[2][:80]


def _pdf_page_sections(page: fitz.Page, page_number: int, start_index: int, exclude_rects: list | None = None, previous_heading: str | None = None) -> list[dict]:
    if exclude_rects:
        raw_lines = []
        for block in page.get_text("dict").get("blocks", []):
            for line in block.get("lines", []):
                box = fitz.Rect(line["bbox"])
                center = (box.tl + box.br) / 2
                if not any(fitz.Rect(rect).contains(center) for rect in exclude_rects):
                    raw_lines.append("".join(span.get("text", "") for span in line.get("spans", [])))
    else:
        raw_lines = page.get_text().splitlines()
    lines = [
        cleaned for line in raw_lines
        if (cleaned := _clean_caption(line)) and not _is_page_marker(cleaned)
    ]
    visual_headings = _pdf_visual_headings(page)
    heading_re = re.compile(
        r"^(?:第\s*[一二三四五六七八九十\d]+\s*[章节]\s*)?"
        r"(?:\d+(?:\.\d+){1,3}\s+)?"
        r"(摘要|绪论|引言|背景|研究现状|研究目标|总体框架|系统架构|技术路线|"
        r"数据集|实验设置|实验结果|结果分析|消融实验|误差分析|系统实现|"
        r"创新点|局限性|总结|结论|未来工作|[\u4e00-\u9fffA-Za-z][^。！？]{2,55})$",
        re.IGNORECASE,
    )
    markers: list[tuple[int, str]] = []
    for index, line in enumerate(lines):
        if not line or len(line) > 85 or line.startswith(("图", "表", "Figure", "Table")):
            continue
        numbered = re.match(r"^[1-9]\d?(?:\.\d{1,2}){1,3}\s+[^。！？]{2,70}$", line)
        named = re.match(r"^(?:摘\s*要|Abstract|第\s*[一二三四五六七八九十\d]+\s*章|结论与展望|总结与展望|参考文献|致谢)$", line, re.IGNORECASE)
        if numbered or named or heading_re.match(line) or (
            line in visual_headings
            and re.match(
                r"^(?:第\s*[一二三四五六七八九十\d]+\s*[章节篇部]|[（(][一二三四五六七八九十\d]+[）)]|[1-9]\d?[.、])",
                line,
            )
        ):
            if re.search(r"https?://|www\.|doi", line, re.IGNORECASE):
                continue
            markers.append((index, line))
    # 正文中大量短行不能都当标题；仅保留明确编号/章节标题。
    strong_pattern = re.compile(
        r"^(?:"
        r"[1-9]\d?(?:\.\d{1,2}){1,3}[.、]?\s*[^。！？]{2,70}|"
        r"[1-9]\d?[.、]\s*[^。！？]{2,55}|"
        r"[（(][一二三四五六七八九十\d]+[）)]\s*[^。！？]{2,55}|"
        r"摘\s*要|Abstract|"
        r"第\s*[一二三四五六七八九十\d]+\s*[章节篇部]\s*[^。！？]{0,55}|"
        r"结论与展望|总结与展望|参考文献|致谢"
        r")$",
        re.IGNORECASE,
    )
    strong = [
        item for item in markers
        if strong_pattern.match(item[1]) or item[1] in visual_headings
    ]
    markers = strong
    if not markers:
        return [_section(start_index, previous_heading or _pdf_page_title(page, page_number), "\n".join(lines), page_number)]
    result = []
    prefix = "\n".join(lines[:markers[0][0]]).strip()
    # Text before the first heading continues the previous section. Attaching it
    # to the next heading swaps table/task meanings at page boundaries.
    if prefix and len(re.sub(r"\s|\d", "", prefix)) >= 8:
        result.append(_section(start_index, previous_heading or _pdf_page_title(page, page_number), prefix, page_number))
    for marker_index, (line_index, heading) in enumerate(markers):
        end = markers[marker_index + 1][0] if marker_index + 1 < len(markers) else len(lines)
        body = "\n".join(lines[line_index + 1:end]).strip()
        result.append(_section(start_index + len(result), heading, body or heading, page_number))
    return result


def _semantic_role(title: str, text: str) -> str:
    for role, pattern in SEMANTIC_PATTERNS:
        if pattern.search(title or ""):
            return role
    sample = (text or "")[:1600]
    for role, pattern in SEMANTIC_PATTERNS:
        if pattern.search(sample):
            return role
    return "content"


def _enrich_sections(sections: list[dict]) -> list[dict]:
    for section in sections:
        text = section.get("text", "")
        title = section.get("title", "")
        role = section.get("semanticRole") or _semantic_role(title, text)
        has_metrics = bool(re.search(r"\d+(?:\.\d+)?\s*(?:%|倍|ms|毫秒|秒|分钟|小时|MB|GB)|Accuracy|Precision|Recall|F1", text, re.IGNORECASE))
        section["semanticRole"] = role
        section["evidenceTypes"] = list(dict.fromkeys([
            *(["metric"] if has_metrics else []),
            *(["method"] if role in {"method", "architecture"} else []),
            *(["finding"] if role in {"data", "insight", "conclusion"} else []),
            "claim",
        ]))
        section["importance"] = round(min(1.0, 0.35 + (0.22 if has_metrics else 0) + (0.18 if role in {"method", "architecture", "data", "conclusion"} else 0) + min(len(text), 2200) / 10000), 3)
    return sections


def _document_profile(name: str, sections: list[dict], figures: list[dict], tables: list[dict]) -> dict:
    suffix = Path(name).suffix.lower()
    sample = "\n".join(
        [str(name), *[f"{section.get('title', '')}\n{section.get('text', '')[:500]}" for section in sections[:40]]]
    )
    business_signals = sum(
        bool(re.search(pattern, sample, re.IGNORECASE))
        for pattern in (
            r"商业计划书|创业计划书|business\s+plan",
            r"市场分析|行业分析|商机分析",
            r"产品概况|产品服务|解决方案",
            r"市场营销|竞争环境|竞品",
            r"财务预测|现金流量|销售预测",
            r"资本结构|融资|资金分配",
            r"团队结构|公司管理|组织架构",
            r"风险分析|风险对冲",
        )
    )
    presentation_mode = "business-plan" if business_signals >= 3 else "academic-report"
    kind = (
        "商业计划书" if presentation_mode == "business-plan"
        else "论文/报告" if suffix in {".pdf", ".docx"}
        else "数据表" if suffix in {".xlsx", ".xlsm", ".csv"}
        else "既有演示" if suffix == ".pptx"
        else "视觉材料" if suffix in {".png", ".jpg", ".jpeg", ".webp"}
        else "文本材料"
    )
    coverage = Counter(section.get("semanticRole", "content") for section in sections)
    return {
        "kind": kind,
        "presentationMode": presentation_mode,
        "semanticCoverage": dict(coverage),
        "outline": [{"section": item["id"], "title": item.get("title", ""), "page": item.get("page"), "role": item.get("semanticRole"),
                     "headingPath": item.get("headingPath", []), "mainPoint": item.get("understanding", {}).get("mainPoint", "")} for item in sections],
        "visualEvidence": {"figures": len(figures), "tables": len(tables)},
    }


def _parse_text_sections(name: str, text: str) -> list[dict]:
    heading_re = re.compile(r"(?m)^(#{1,4})\s+(.+?)\s*$|^(\d+(?:\.\d+){0,3})\s+([^\n]{2,90})$")
    matches = list(heading_re.finditer(text))
    if not matches:
        chunks = [item for item in re.split(r"\n\s*\n+", text) if item.strip()]
        return [_section(index, Path(name).stem if index == 1 else f"第 {index} 节", chunk) for index, chunk in enumerate(chunks or [text], 1)]
    result = []
    if text[:matches[0].start()].strip():
        result.append(_section(1, Path(name).stem, text[:matches[0].start()]))
    ancestry = []
    for match_index, match in enumerate(matches):
        start = match.end()
        end = matches[match_index + 1].start() if match_index + 1 < len(matches) else len(text)
        title = (match.group(2) or f"{match.group(3)} {match.group(4)}").strip()
        body = text[start:end].strip()
        level = len(match.group(1)) if match.group(1) else len(match.group(3).split("."))
        while ancestry and ancestry[-1][0] >= level:
            ancestry.pop()
        ancestry.append((level, title))
        result.append({**_section(len(result) + 1, title, body), "headingLevel": level,
                       "headingPath": [item[1] for item in ancestry]})
    return result or [_section(1, Path(name).stem, text)]


def _caption_blocks(page: fitz.Page, pattern: re.Pattern[str]) -> list[dict]:
    matches = []
    for block in page.get_text("blocks"):
        raw_text = str(block[4])
        for line in raw_text.splitlines() or [raw_text]:
            line = _clean_caption(line)
            match = pattern.match(line)
            if match:
                matches.append({
                    "text": _clean_caption(f"{match.group(1)} {match.group(2)}"),
                    "bbox": [round(float(value), 2) for value in block[:4]],
                })
                break
    return matches


def _nearest_caption(image_bbox: fitz.Rect, captions: list[dict]) -> str:
    if not captions:
        return ""
    def distance(item: dict) -> tuple[float, float]:
        box = fitz.Rect(item["bbox"])
        vertical = box.y0 - image_bbox.y1 if box.y0 >= image_bbox.y1 else image_bbox.y0 - box.y1
        horizontal = max(0.0, max(image_bbox.x0, box.x0) - min(image_bbox.x1, box.x1))
        below_penalty = 0.0 if box.y0 >= image_bbox.y1 else 45.0
        return (abs(vertical) + horizontal * 0.25 + below_penalty, abs(box.x0 - image_bbox.x0))
    return min(captions, key=distance)["text"]


def _figure_kind(caption: str, page_text: str) -> str:
    caption_text = caption.lower()
    if re.search(r"prompt|提示词|指令", caption_text, re.IGNORECASE):
        return "prompt"
    if re.search(r"f1|accuracy|precision|recall|性能|分布|分类|标签|成本|耗时|时间|效率|完整性|支撑率|结果|指标|得分", caption_text, re.IGNORECASE):
        return "chart"
    if re.search(r"architecture|framework|workflow|pipeline|架构|框架|流程|技术路线|向量库|rag|检索|一致性|合规", caption_text, re.IGNORECASE):
        return "diagram"
    text = page_text[:1200].lower()
    if not caption and re.search(r"f1|accuracy|precision|recall|性能|分布|成本|耗时|效率|完整性|支撑率|结果", text, re.IGNORECASE):
        return "chart"
    if not caption and re.search(r"架构|框架|流程|技术路线|向量库|rag|检索|一致性|合规", text, re.IGNORECASE):
        return "diagram"
    return "figure"


def _extract_pdf_figures(
    page: fitz.Page,
    page_number: int,
    name: str,
    asset_dir: Path | None,
) -> list[dict]:
    page_text = page.get_text()
    captions = _caption_blocks(page, FIGURE_CAPTION_RE)
    page_area = max(1.0, float(page.rect.width * page.rect.height))
    figures: list[dict] = []
    seen: set[tuple[int, int, int, int, int]] = set()
    for image_index, info in enumerate(page.get_image_info(xrefs=True), 1):
        bbox = fitz.Rect(info.get("bbox"))
        key = (
            int(info.get("xref") or 0),
            round(bbox.x0), round(bbox.y0), round(bbox.x1), round(bbox.y1),
        )
        if key in seen:
            continue
        seen.add(key)
        pixel_width = int(info.get("width") or 0)
        pixel_height = int(info.get("height") or 0)
        if (
            pixel_width < 100 or pixel_height < 70
            or bbox.width < 48 or bbox.height < 36
            or (bbox.width * bbox.height) / page_area < 0.008
        ):
            continue
        caption = _nearest_caption(bbox, captions)
        path = None
        rendered_width, rendered_height = pixel_width, pixel_height
        if asset_dir is not None:
            asset_dir.mkdir(parents=True, exist_ok=True)
            target = asset_dir / f"page-{page_number:03d}-figure-{image_index:02d}.png"
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), clip=bbox, alpha=False)
            pixmap.save(target)
            rendered_width, rendered_height = pixmap.width, pixmap.height
            path = str(target.resolve())
        kind = _figure_kind(caption, page_text)
        figures.append({
            "id": f"FIG-{page_number:03d}-{image_index:02d}",
            "document": name,
            "page": page_number,
            "xref": int(info.get("xref") or 0),
            "bbox": [round(float(value), 2) for value in bbox],
            "width": rendered_width,
            "height": rendered_height,
            "aspectRatio": round(rendered_width / max(1, rendered_height), 4),
            "path": path,
            "caption": caption,
            "kind": kind,
            "provenance": f"{name} · 第 {page_number} 页" + (f" · {caption}" if caption else ""),
        })
    return figures


def _decode_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _append_tabular_sections(name: str, title: str, rows: list, sections: list, tables: list) -> None:
    rows = [row for row in rows if any(value is not None and str(value).strip() for value in row)]
    if not rows:
        return
    for start in range(1, max(2, len(rows)), 50):
        section_id = f"S{len(sections) + 1:03d}"
        table = table_record([rows[0], *rows[start:start + 50]], table_id=f"TABLE-DATA-{len(tables) + 1:03d}",
                             document=name, section=section_id, caption=title, row_offset=start - 1)
        table["name"] = title
        tables.append(table)
        sections.append({**_section(len(sections) + 1, title, table_text(table)),
                         "semanticRole": "data", "headingPath": [title],
                         "blocks": [{"type": "table", "tableId": table["id"], "text": table_text(table)}]})


def _extract_markdown_tables(name: str, sections: list, tables: list) -> None:
    for section in sections:
        lines = section["text"].splitlines()
        output, index = [], 0
        while index < len(lines):
            if index + 1 < len(lines) and "|" in lines[index] and re.fullmatch(r"\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*", lines[index + 1]):
                rows = [[cell.strip() for cell in lines[index].strip().strip("|").split("|")]]
                index += 2
                while index < len(lines) and "|" in lines[index] and lines[index].strip():
                    rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                    index += 1
                table = table_record(rows, table_id=f"TABLE-MD-{len(tables) + 1:03d}", document=name,
                                     section=section["id"], caption=section["title"])
                tables.append(table)
                output.extend(["", table_text(table), ""])
            else:
                output.append(lines[index])
                index += 1
        section["text"] = "\n".join(output)


def parse_source(
    name: str,
    data: bytes,
    media_type: str = "application/octet-stream",
    asset_dir: Path | None = None,
) -> dict:
    suffix = Path(name).suffix.lower()
    sections, figures, tables, warnings = [], [], [], []
    if suffix == ".pdf":
        pdf = fitz.open(stream=data, filetype="pdf")
        for index, page in enumerate(pdf, 1):
            page_text = page.get_text()
            if len(re.sub(r"\s", "", page_text)) < 20:
                warnings.append({"code": "ocr-required", "page": index, "message": f"第 {index} 页缺少可读取文字，需要文字识别或补充文本"})
            figures.extend(_extract_pdf_figures(page, index, name, asset_dir))
            captions = _caption_blocks(page, TABLE_CAPTION_RE)
            try:
                detected = page.find_tables().tables
            except (ValueError, RuntimeError) as exc:
                detected = []
                warnings.append({"code": "table-extraction-failed", "page": index, "message": str(exc)[:180]})
            if not detected and captions:
                detected = _aligned_metric_tables(page, captions)
            if _is_contents_page(page_text):
                warnings.append({"code": "contents-page-skipped", "page": index, "message": f"第 {index} 页为目录，已从正文证据中排除"})
                continue
            previous_section = sections[-1] if sections else None
            page_sections = _pdf_page_sections(page, index, len(sections) + 1, [table.bbox for table in detected], previous_section["title"] if previous_section else None)
            sections.extend(page_sections)
            for detected_table in detected:
                rows = detected_table.extract()
                if not rows:
                    continue
                if detected_table.header.external:
                    rows.insert(0, detected_table.header.names)
                table_y = detected_table.bbox[1]
                # A table can continue from the previous page before the next
                # chapter heading.  Defaulting to the first new section made
                # those continuation tables appear under the upcoming chapter
                # (for example, a team-management table became "财务预测").
                owner = previous_section or page_sections[0]
                for section in page_sections:
                    locations = page.search_for(section["title"])
                    if locations and min(location.y0 for location in locations) <= table_y:
                        owner = section
                caption = _nearest_caption(fitz.Rect(detected_table.bbox), captions)
                table = table_record(rows, table_id=f"TABLE-{index:03d}-{len(tables) + 1:02d}",
                                     document=name, section=owner["id"], page=index, caption=caption or owner["title"])
                table["bbox"] = list(detected_table.bbox)
                tables.append(table)
                owner["text"] += "\n\n" + table_text(table)
            if captions and not detected:
                warnings.append({"code": "table-text-only", "page": index, "message": "发现表题，但无法可靠恢复行列关系；保留原文，不猜测单元格"})
        pdf.close()
    elif suffix == ".docx":
        sections, figures, tables = parse_docx_structure(name, data, asset_dir)
    elif suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        for sheet in workbook.worksheets:
            rows = []
            missing_formula_results = 0
            for index, (row, formula_row) in enumerate(zip(sheet.iter_rows(), formulas[sheet.title].iter_rows())):
                if index >= 10000:
                    warnings.append({"code": "row-limit", "message": f"{sheet.title} 超过 10000 行，后续行未读取，请按主题拆分"})
                    break
                missing_formula_results += sum(cell.value is None and formula.data_type == "f" for cell, formula in zip(row, formula_row))
                values = []
                for cell in row:
                    number_format = re.sub(r'"[^"]*"|\\.', "", cell.number_format or "")
                    values.append(f"{cell.value * 100:.12g}%" if type(cell.value) in {int, float} and "%" in number_format else cell.value)
                rows.append(values)
            if missing_formula_results:
                warnings.append({"code": "formula-result-unavailable", "message": f"{sheet.title} 有 {missing_formula_results} 个公式缺少已计算结果，暂记为未提供；请在表格软件中计算并保存后重新上传"})
            _append_tabular_sections(name, sheet.title, rows, sections, tables)
        workbook.close()
        formulas.close()
    elif suffix == ".csv":
        text = _decode_text(data)
        rows = []
        for index, row in enumerate(csv.reader(io.StringIO(text))):
            if index >= 10000:
                warnings.append({"code": "row-limit", "message": "数据超过 10000 行，后续行未读取，请按主题拆分"})
                break
            rows.append(row)
        _append_tabular_sections(name, Path(name).stem, rows, sections, tables)
    elif suffix == ".pptx":
        for index, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
            paragraphs = []
            def read_shapes(shapes, page_number=index, output=paragraphs):
                for shape in shapes:
                    if hasattr(shape, "shapes"):
                        read_shapes(shape.shapes, page_number, output)
                    elif getattr(shape, "has_table", False):
                        table = table_record([[cell.text for cell in row.cells] for row in shape.table.rows],
                                             table_id=f"TABLE-PPTX-{page_number}-{len(tables) + 1}", document=name, section=f"S{page_number:03d}", page=page_number)
                        tables.append(table)
                        output.append(table_text(table))
                    elif hasattr(shape, "text"):
                        output.append(shape.text)
            read_shapes(slide.shapes)
            title = slide.shapes.title.text if slide.shapes.title else f"第 {index} 页幻灯片"
            sections.append(_section(index, title, "\n\n".join(paragraphs), index))
    elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        image = Image.open(io.BytesIO(data))
        figures.append(
            {"name": name, "width": image.width, "height": image.height, "mode": image.mode}
        )
        sections.append(_section(1, name, f"图片尺寸 {image.width}×{image.height}"))
        warnings.append({"code": "ocr-required", "message": "已读取图片尺寸；图中文字与数据尚未识别，请补充文本或可读取的文档"})
    elif suffix in {".html", ".htm"}:
        soup = BeautifulSoup(data, "html.parser")
        for node in soup(["script", "style"]):
            node.decompose()
        sections.append(
            _section(1, soup.title.string if soup.title else name, soup.get_text("\n", strip=True))
        )
    else:
        text = _decode_text(data)
        sections = _parse_text_sections(name, text)
        _extract_markdown_tables(name, sections, tables)
    if not any(section.get("text", "").strip() for section in sections) and not warnings:
        warnings.append({"code": "empty-document", "message": "没有提取到正文，请检查文件内容或补充可读取的文本"})
    if sections and _is_page_marker(sections[0].get("title", "")):
        sections[0]["title"] = Path(name).stem
    for section in sections:
        section["tableIds"] = [table["id"] for table in tables if table.get("section") == section["id"]]
        section["figureIds"] = [figure["id"] for figure in figures if figure.get("section") == section["id"] or (section.get("page") and figure.get("page") == section["page"])]
        section["understanding"] = section_understanding(section)
    sections = _enrich_sections(sections)
    return {
        "parserVersion": PARSER_VERSION,
        "name": name,
        "mediaType": media_type,
        "sha256": hashlib.sha256(data).hexdigest(),
        "sections": sections,
        "figures": figures,
        "tables": tables,
        "documentProfile": _document_profile(name, sections, figures, tables),
        "readingWarnings": warnings,
        "stats": {
            "sections": len(sections),
            "figures": len(figures),
            "tables": len(tables),
            "characters": sum(len(x["text"]) for x in sections),
        },
    }
