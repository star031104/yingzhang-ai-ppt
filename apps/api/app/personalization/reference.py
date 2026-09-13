"""Extract data-only visual preferences from OOXML. Do not retain slide text."""

import io
import json
import zipfile
from collections import Counter
from xml.etree import ElementTree as ET

from pptx import Presentation

from app.personalization.schemas import OPTIONS, validate_preference

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"a": A}


def extract_reference(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if (
            len(archive.infolist()) > 5000
            or sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024
        ):
            raise ValueError("参考稿解压后过大")
        theme_names = sorted(
            name
            for name in archive.namelist()
            if name.startswith("ppt/theme/theme") and name.endswith(".xml")
        )
        theme = ET.fromstring(archive.read(theme_names[0])) if theme_names else None
    theme_colors, theme_fonts = {}, []
    if theme is not None:
        for element in theme.findall(".//a:clrScheme/*", NS):
            child = next(iter(element), None)
            if child is not None:
                color = child.get("val") if child.tag.endswith("srgbClr") else child.get("lastClr")
                if color and len(color) == 6:
                    theme_colors[element.tag.rsplit("}", 1)[-1]] = "#" + color
        theme_fonts = [item.get("typeface") for item in theme.findall(".//a:fontScheme//a:ea", NS)]
        theme_fonts += [
            item.get("typeface")
            for item in theme.findall(".//a:fontScheme//a:font", NS)
            if item.get("script") == "Hans"
        ]
        theme_fonts += [
            item.get("typeface") for item in theme.findall(".//a:fontScheme//a:latin", NS)
        ]
    deck = Presentation(io.BytesIO(data))
    if len(deck.slides) > 100:
        raise ValueError("参考稿最多支持 100 页")
    fonts, backgrounds, foregrounds, fills, variants = (
        Counter(),
        Counter(),
        Counter(),
        Counter(),
        {},
    )

    def color_of(color):
        try:
            from pptx.enum.dml import MSO_COLOR_TYPE

            if color.type == MSO_COLOR_TYPE.RGB:
                return "#" + str(color.rgb)
        except (AttributeError, TypeError, ValueError):
            pass
        return None

    def visit(shapes, counts):
        for shape in shapes:
            if hasattr(shape, "shapes"):
                visit(shape.shapes, counts)
            if getattr(shape, "has_chart", False):
                counts["chart"] += 1
            if getattr(shape, "has_table", False):
                counts["table"] += 1
            try:
                color = color_of(shape.fill.fore_color)
                if color:
                    fills[color] += max(1, shape.width * shape.height)
            except (AttributeError, TypeError, ValueError):
                pass
            if getattr(shape, "has_text_frame", False):
                counts["text"] += 1
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name in OPTIONS["font_family"]:
                            fonts[run.font.name] += len(run.text)
                        color = color_of(run.font.color)
                        if color:
                            foregrounds[color] += len(run.text)

    for index, slide in enumerate(deck.slides):
        counts = Counter()
        visit(slide.shapes, counts)
        try:
            background = color_of(slide.background.fill.fore_color)
            if background:
                backgrounds[background] += 1
        except (AttributeError, TypeError, ValueError):
            pass
        if index == 0 and counts["text"] <= 4:
            variants["cover"] = "minimal-cover"
        if counts["table"]:
            variants["data"] = "table-highlight"
        elif counts["chart"] and "data" not in variants:
            variants["data"] = "chart-focus"
        elif counts["text"] == 2:
            variants.setdefault("content", "statement")
        elif counts["text"] in {3, 4}:
            variants.setdefault("content", "split")
    font = (
        fonts.most_common(1)[0][0]
        if fonts
        else next((value for value in theme_fonts if value in OPTIONS["font_family"]), None)
    )
    ink = foregrounds.most_common(1)[0][0] if foregrounds else theme_colors.get("dk1", "#182033")
    bg = backgrounds.most_common(1)[0][0] if backgrounds else theme_colors.get("lt1", "#FFFFFF")
    accents = [color for color, _ in fills.most_common() if color not in {bg, ink}]
    palette = {
        "bg": bg,
        "ink": ink,
        "primary": accents[0] if accents else theme_colors.get("accent1", "#625BF6"),
        "accent": accents[1] if len(accents) > 1 else theme_colors.get("accent2", "#78E3C5"),
        "deep": theme_colors.get("dk2", "#20284D"),
        "line": theme_colors.get("lt2", "#E3E7F0"),
    }
    style = {"palette": palette, "preferredByRole": variants}
    if font:
        style["fontFamily"] = font
    from app.personalization.reference_layouts import layout_tokens
    inherited = layout_tokens(data)
    style["layouts"] = inherited["layouts"]
    value = validate_preference("reference_style", json.dumps(style))
    return value, {
        "slideCount": len(deck.slides),
        "inheritance": inherited,
        "supportedFont": font,
        "themeRead": theme is not None,
        "retainedOriginal": False,
        "limitations": [
            "复杂母版、动画和逐对象位置不转换成通用规则",
            "参考颜色与页型为待确认建议；生成时仍与普通样式对照",
            "字体在目标设备的可用性需要核对",
        ],
    }
