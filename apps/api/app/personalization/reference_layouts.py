"""Resolve slide/layout/master font inheritance into bounded layout tokens."""
import io
from collections import Counter, defaultdict

from pptx import Presentation

from app.personalization.schemas import OPTIONS


def layout_tokens(data):
    deck = Presentation(io.BytesIO(data))
    collected = defaultdict(list)
    warnings = set()
    for index, slide in enumerate(deck.slides):
        text_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
        role = "cover" if index == 0 else "content"
        if any(getattr(shape, "has_chart", False) or getattr(shape, "has_table", False) for shape in slide.shapes):
            role = "data"
        if any(hasattr(shape, "shapes") for shape in slide.shapes):
            warnings.add("组合对象保留在原生参考模板中，不展平为通用文本布局")
        boxes = []
        for shape in sorted(text_shapes, key=lambda item: (item.top, item.left))[:12]:
            fonts = Counter()
            sizes = []
            candidates = [shape]
            if shape.is_placeholder:
                idx = shape.placeholder_format.idx
                candidates += [item for item in slide.slide_layout.placeholders if item.placeholder_format.idx == idx]
                candidates += [item for item in slide.slide_layout.slide_master.placeholders if item.placeholder_format.type == shape.placeholder_format.type]
            for candidate in candidates:
                for paragraph in candidate.text_frame.paragraphs:
                    if paragraph.font.size:
                        sizes.append(paragraph.font.size.pt)
                    if paragraph.font.name in OPTIONS["font_family"]:
                        fonts[paragraph.font.name] += 1
                    for run in paragraph.runs:
                        if run.font.size:
                            sizes.append(run.font.size.pt)
                        if run.font.name in OPTIONS["font_family"]:
                            fonts[run.font.name] += max(1, len(run.text))
                if sizes:
                    break
            x, y = float(shape.left / deck.slide_width), float(shape.top / deck.slide_height)
            w, h = float(shape.width / deck.slide_width), float(shape.height / deck.slide_height)
            if x < 0 or y < 0 or w <= 0 or h <= 0 or x + w > 1.005 or y + h > 1.005:
                warnings.add("越界对象不作为通用布局规则")
                continue
            box = {"x": round(x, 5), "y": round(y, 5), "w": round(min(w, 1-x), 5), "h": round(min(h, 1-y), 5),
                   "fontSize": round(max(14, min(64, max(sizes) * 1280 / (deck.slide_width / 914400 * 72))) if sizes else (44 if not boxes else 22), 1)}
            if fonts:
                box["fontFamily"] = fonts.most_common(1)[0][0]
            boxes.append(box)
        if 2 <= len(boxes) <= 8 and role != "data":
            collected[role].append(boxes)
    layouts = {}
    for role, examples in collected.items():
        common_size = Counter(len(boxes) for boxes in examples).most_common(1)[0][0]
        layouts[role] = next(boxes for boxes in examples if len(boxes) == common_size)
    return {"layouts": layouts, "masterCount": len(deck.slide_masters),
            "layoutCount": sum(len(master.slide_layouts) for master in deck.slide_masters),
            "themeInheritance": True, "warnings": sorted(warnings)}
