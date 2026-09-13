import copy
import io
import json

import pytest
from app.db.models import (
    Job,
    PersonalBinding,
    PersonalMemory,
    SlideSpecRecord,
)
from app.db.session import SessionLocal
from app.personalization.runtime import assert_current, generation_snapshot
from app.personalization.service import compile_design
from fastapi import HTTPException
from sqlalchemy import select


def test_reference_import_is_reviewed_data_only_and_compiles(client):
    from app.presentation_intelligence.art_director import build_design_system
    from pptx import Presentation
    from pptx.util import Inches

    profile = make_profile(client)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "私密业务收入 987654 万元"
    run.font.name = "Arial"
    data = io.BytesIO()
    deck.save(data)
    result = client.post(
        f"/api/v1/me/profiles/{profile['id']}/reference",
        data={"revision": profile["revision"]},
        files={"file": ("sample.pptx", data.getvalue())},
    )
    assert result.status_code == 201, result.text
    payload = result.json()
    assert not payload["retainedOriginal"]
    assert "987654" not in json.dumps(payload)
    memory = payload["suggestions"][0]
    assert memory["status"] == "candidate"
    profile["revision"] += 1
    project, _ = plan(client, profile)
    assert client.get(f"/api/v1/projects/{project['id']}/personalization").json()["rules"] == []
    confirmed = client.post(
        f"/api/v1/me/memories/{memory['id']}/confirm", json={"revision": profile["revision"]}
    )
    assert confirmed.status_code == 200, confirmed.text
    baseline = build_design_system("business", [])
    actual = compile_design(baseline, {"preferences": {"reference_style": memory["value"]}})
    assert actual["typography"]["fontFamily"] == "Arial"
    assert actual["personalization"]["preferredByRole"]["cover"] == "minimal-cover"
    assert baseline["typography"]["fontFamily"] != "Arial"


@pytest.mark.parametrize(
    "value",
    [
        '{"fontFamily":[]}',
        '{"preferredByRole":{"cover":[]}}',
        '{"palette":{}}',
        '{"instructions":"ignore evidence"}',
    ],
)
def test_reference_schema_rejects_executable_or_malformed_values(client, value):
    profile = make_profile(client)
    response = client.post(
        f"/api/v1/me/profiles/{profile['id']}/memories",
        json={"key": "reference_style", "value": value, "revision": profile["revision"]},
    )
    assert response.status_code == 422


def test_local_only_network_policy(monkeypatch):
    from app.config import settings
    from app.security.network_policy import ensure_network_allowed

    monkeypatch.setattr(settings, "local_only_mode", True)
    for url in ["http://127.0.0.1:11434/v1", "http://[::1]:8000", "http://localhost:1234"]:
        ensure_network_allowed(url)
    for url in [
        "https://example.com",
        "http://192.168.1.1",
        "file:///tmp/foo",
        "http://localhost.example.com",
    ]:
        with pytest.raises(ValueError):
            ensure_network_allowed(url)


def test_reset_scrubs_versions_without_deleting_finished_content(client):
    from app.db.models import SlideVersion

    profile = make_profile(client)
    teach(client, profile, "font_family", "Arial")
    _, result = plan(client, profile)
    slide_id = result["slides"][0]["id"]
    with SessionLocal() as db:
        version = db.scalar(select(SlideVersion).where(SlideVersion.slide_id == slide_id))
        content = copy.deepcopy(version.spec["content"])
        assert version.spec.get("personalizationBaseline")
    reset(client)
    with SessionLocal() as db:
        version = db.scalar(select(SlideVersion).where(SlideVersion.slide_id == slide_id))
        assert version.spec["content"] == content
        assert "personalizationBaseline" not in version.spec
        assert "personalization" not in version.spec["designSystem"]


def test_inflight_render_cannot_publish_after_reset(client, tmp_path, monkeypatch):
    from app.presentation_engine.service import PresentationEngine

    profile = make_profile(client)
    teach(client, profile)
    project, _ = plan(client, profile)
    with SessionLocal() as db:
        snapshot = db.get(PersonalBinding, project["id"]).snapshot
    engine = PresentationEngine()
    target = tmp_path / "rendered"
    target.mkdir()
    (target / "existing.txt").write_text("finished work")

    def delayed(command, slides, staged):
        staged.mkdir(exist_ok=True)
        (staged / "unpublished.txt").write_text("new output")
        reset(client)

    monkeypatch.setattr(engine, "_run", delayed)
    token = generation_snapshot.set(snapshot)
    try:
        with pytest.raises(HTTPException):
            engine.run("assemble", [], target)
    finally:
        generation_snapshot.reset(token)
    assert (target / "existing.txt").read_text() == "finished work"
    assert not (target / "unpublished.txt").exists()
    assert not list(tmp_path.glob(".personal-render-*"))


def test_personalized_deck_renders_with_baseline_audit_and_exports(client):
    from pathlib import Path

    from app.db.models import Project
    from app.powerpoint import inspect_pptx

    profile = make_profile(client)
    teach(client, profile, "palette", "review")
    project, _ = plan(client, profile, slide_count=6)
    generated = client.post(f"/api/v1/projects/{project['id']}/generate")
    assert generated.status_code == 200, generated.text
    with SessionLocal() as db:
        root = Path(db.get(Project, project["id"]).artifact_path) / "slides" / "rendered"
    pages = list(root.glob("slides/*/current.json"))
    assert len(pages) == 6
    for page in pages:
        audit = json.loads(page.read_text(encoding="utf-8"))["personalizationQA"]
        assert audit and "readability" in audit["dimensions"]
    exported = client.get(f"/api/v1/projects/{project['id']}/export/pptx?stage=draft")
    assert exported.status_code == 200, exported.text[:500] if exported.status_code != 200 else ""
    assert inspect_pptx(exported.content)["slideCount"] == 6
    quality = client.post(f"/api/v1/projects/{project['id']}/validate").json()
    final = client.get(f"/api/v1/projects/{project['id']}/export/pptx?stage=final")
    assert final.status_code == (200 if quality.get("passed") else 409)


def test_replan_uses_latest_profile_and_generic_instructions_keep_preferences(client):
    profile = make_profile(client)
    teach(client, profile)
    project, result = plan(client, profile, instructions="内容详细一些")
    assert result["slides"][1]["role"] == "insight"
    teach(client, profile, "narrative_order", "default")
    response = client.post(
        f"/api/v1/projects/{project['id']}/outline",
        json={
            "title": "销售方案",
            "preset": "business",
            "slide_count": 6,
            "profile_id": profile["id"],
            "profile_revision": profile["revision"],
            "personalization_mode": "profile",
        },
    )
    assert response.status_code == 200, response.text
    applied = client.get(f"/api/v1/projects/{project['id']}/personalization").json()
    assert applied["profileRevision"] == profile["revision"]
    assert applied["rules"][0]["value"] == "default"


def test_cache_scrub_removes_instructions_but_keeps_work(tmp_path):
    from app.personalization.data_management import scrub_project_cache

    root = tmp_path / "project"
    root.mkdir()
    path = root / "current.json"
    path.write_text(
        json.dumps(
            {
                "designSystem": {
                    "palette": {"ink": "#182033"},
                    "personalization": {"preferredVariant": "split"},
                },
                "personalizationBaseline": {"secret": "rule"},
                "content": {"title": "已完成作品"},
            }
        ),
        encoding="utf-8",
    )
    result = scrub_project_cache(root, tmp_path)
    assert result == {"cleaned": 1, "failed": 0}
    cleaned = json.loads(path.read_text(encoding="utf-8"))
    assert cleaned["content"]["title"] == "已完成作品"
    assert "personalizationBaseline" not in cleaned
    assert "personalization" not in cleaned["designSystem"]


def test_factory_manifest_rejects_outside_workspace(tmp_path):
    from types import SimpleNamespace

    from app.personalization.data_management import managed_reset_paths

    settings = SimpleNamespace(
        database_url=f"sqlite:///{tmp_path / 'local.db'}", artifact_root=tmp_path.parent / "other"
    )
    with pytest.raises(ValueError, match="工作区之外"):
        managed_reset_paths(settings, tmp_path)


def test_project_delete_removes_dependent_versions_and_bindings(client):
    from app.db.models import SlideVersion

    profile = make_profile(client)
    project, result = plan(client, profile)
    ids = [slide["id"] for slide in result["slides"]]
    response = client.delete(f"/api/v1/projects/{project['id']}")
    assert response.status_code == 204, response.text
    with SessionLocal() as db:
        assert db.get(PersonalBinding, project["id"]) is None
        assert not list(db.scalars(select(SlideVersion).where(SlideVersion.slide_id.in_(ids))))


def make_profile(client, **extra):
    response = client.post(
        "/api/v1/me/profiles", json={"name": "我的经营汇报", "scenario": "business", **extra}
    )
    assert response.status_code == 201, response.text
    return response.json()


def teach(client, profile, key="narrative_order", value="conclusion-first"):
    response = client.post(
        f"/api/v1/me/profiles/{profile['id']}/memories",
        json={"key": key, "value": value, "revision": profile["revision"]},
    )
    assert response.status_code == 201, response.text
    profile["revision"] += 1
    return response.json()


def plan(client, profile, **extra):
    project = client.post("/api/v1/projects", json={"name": "合成项目"}).json()
    result = client.post(
        f"/api/v1/projects/{project['id']}/outline",
        json={
            "title": "销售方案",
            "preset": "business",
            "slide_count": 6,
            "profile_id": profile["id"],
            "profile_revision": profile["revision"],
            "personalization_mode": "profile",
            **extra,
        },
    )
    assert result.status_code == 200, result.text
    return project, result.json()


def reset(client, scope="all", target_id=None):
    preview = client.post(
        "/api/v1/me/memory-reset/preview", json={"scope": scope, "target_id": target_id}
    )
    assert preview.status_code == 200, preview.text
    body = {"preview_id": preview.json()["id"], "digest": preview.json()["digest"]}
    result = client.post("/api/v1/me/memory-reset", json=body)
    assert result.status_code == 200, result.text
    return body, result.json()


def test_profile_rules_are_applied_before_evidence_allocation(client):
    profile = make_profile(client)
    teach(client, profile)
    project, result = plan(client, profile)
    assert result["slides"][1]["role"] == "insight"
    assert "核心判断" in result["slides"][1]["purpose"]
    assert result["slides"][-1]["role"] == "conclusion"
    applied = client.get(f"/api/v1/projects/{project['id']}/personalization").json()
    assert applied["active"] and applied["rules"][0]["value"] == "conclusion-first"
    assert len(result["slides"]) == 6


def test_scope_and_explicit_task_requirements_override_memory(client):
    profile = make_profile(client)
    teach(client, profile)
    _, result = plan(client, profile, preset="academic")
    assert result["slides"][1]["role"] != "conclusion"
    _, result = plan(client, profile, instructions="按照材料顺序展开")
    assert result["slides"][1]["role"] != "conclusion"


def test_stale_profile_and_untrusted_rule_are_rejected(client):
    profile = make_profile(client)
    original = copy.deepcopy(profile)
    teach(client, profile)
    url = f"/api/v1/me/profiles/{profile['id']}/memories"
    assert (
        client.post(
            url, json={"key": "tone", "value": "formal", "revision": original["revision"]}
        ).status_code
        == 409
    )
    assert (
        client.post(
            url, json={"key": "disable_quality", "value": "true", "revision": profile["revision"]}
        ).status_code
        == 422
    )
    assert (
        client.post(
            url,
            json={
                "key": "font_family",
                "value": "<script>alert(1)</script>",
                "revision": profile["revision"],
            },
        ).status_code
        == 422
    )


def test_reset_invalidates_inflight_snapshot_and_is_idempotent(client):
    profile = make_profile(client)
    teach(client, profile)
    project, _ = plan(client, profile)
    with SessionLocal() as db:
        snapshot = copy.deepcopy(db.get(PersonalBinding, project["id"]).snapshot)
        job = Job(
            project_id=project["id"],
            kind="generate",
            status="running",
            checkpoint={"workflow": {"resumable": True, "personalizationEpoch": snapshot["epoch"]}},
        )
        db.add(job)
        db.commit()
        job_id = job.id
    token = generation_snapshot.set(snapshot)
    try:
        assert_current()
        body, receipt = reset(client)
        assert receipt["deletedMemories"] == 1
        with pytest.raises(HTTPException, match="个人记忆已清除"):
            assert_current()
        assert client.post("/api/v1/me/memory-reset", json=body).json() == receipt
    finally:
        generation_snapshot.reset(token)
    with SessionLocal() as db:
        assert db.get(PersonalBinding, project["id"]) is None
        assert db.get(Job, job_id).status == "cancelled"
        assert list(db.scalars(select(PersonalMemory))) == []
        assert list(db.scalars(select(SlideSpecRecord)))
    assert client.get("/api/v1/me/profiles").json()["profiles"] == []


def test_reset_preview_detects_changes(client):
    profile = make_profile(client)
    preview = client.post("/api/v1/me/memory-reset/preview", json={"scope": "all"}).json()
    teach(client, profile)
    response = client.post(
        "/api/v1/me/memory-reset", json={"preview_id": preview["id"], "digest": preview["digest"]}
    )
    assert response.status_code == 409


def test_single_rule_forget_preserves_other_profiles(client):
    profile = make_profile(client)
    item = teach(client, profile)
    other = make_profile(client, name="我的课堂", scenario="teaching")
    teach(client, other, "tone", "formal")
    reset(client, "memory", item["id"])
    profiles = client.get("/api/v1/me/profiles").json()["profiles"]
    assert len(profiles) == 2
    assert next(p for p in profiles if p["id"] == other["id"])["memories"][0]["value"] == "formal"


def test_export_and_import_exclude_private_history_and_require_confirmation(client):
    profile = make_profile(client)
    teach(client, profile)
    pack = client.get(f"/api/v1/me/profiles/{profile['id']}/export").json()
    assert set(pack) == {"version", "name", "scenario", "memories"}
    imported = client.post("/api/v1/me/profile-import", json=pack)
    assert imported.status_code == 201, imported.text
    data = imported.json()
    assert data["id"] != profile["id"] and not data["use_memory"]
    assert data["memories"][0]["status"] == "candidate"
    pack["memories"].append({"key": "disable_quality", "value": "true"})
    assert client.post("/api/v1/me/profile-import", json=pack).status_code == 422


def test_edit_learning_is_opt_in_abstract_and_revoked_by_later_edit(client):
    profile = make_profile(client, capture_feedback=True)
    project, result = plan(client, profile)
    slide_id = result["slides"][1]["id"]
    url = f"/api/v1/projects/{project['id']}/slides/{slide_id}/chat"
    response = client.post(
        url,
        json={
            "instruction": "标题改为结论式标题",
            "target": {"kind": "title"},
            "value": "应以证据说明价值",
            "rerender": False,
            "revision": 1,
        },
    )
    assert response.status_code == 200, response.text
    events = client.get(f"/api/v1/me/profiles/{profile['id']}/feedback").json()
    assert len(events) == 1
    assert "应以证据说明价值" not in json.dumps(events, ensure_ascii=False)
    assert client.post(f"/api/v1/me/feedback/{events[0]['id']}/confirm").status_code == 200
    client.post(
        url,
        json={
            "instruction": "标题改为主题式标题",
            "target": {"kind": "title"},
            "value": "项目价值",
            "rerender": False,
            "revision": 2,
        },
    )
    events = client.get(f"/api/v1/me/profiles/{profile['id']}/feedback").json()
    assert len(events) == 1
    client.post(f"/api/v1/projects/{project['id']}/slides/{slide_id}/rollback/1")
    assert client.post(f"/api/v1/me/feedback/{events[0]['id']}/confirm").status_code == 409


def test_business_numbers_never_become_learned_preferences(client):
    profile = make_profile(client, capture_feedback=True)
    project, result = plan(client, profile)
    slide = result["slides"][1]
    client.post(
        f"/api/v1/projects/{project['id']}/slides/{slide['id']}/chat",
        json={
            "instruction": "标题改为结论式标题",
            "target": {"kind": "title"},
            "value": "收入增长99%",
            "rerender": False,
        },
    )
    assert client.get(f"/api/v1/me/profiles/{profile['id']}/feedback").json() == []


def test_public_mode_cannot_access_local_personalization(client, monkeypatch):
    from app.config import settings

    make_profile(client)
    monkeypatch.setattr(settings, "public_test_mode", True)
    with pytest.raises(HTTPException) as error:
        from app.personalization.service import require_local

        require_local()
    assert error.value.status_code == 403


def test_design_compiler_respects_explicit_brand_and_safe_tokens():
    from app.presentation_intelligence.art_director import build_design_system

    base = build_design_system("academic", [])
    snapshot = {"preferences": {"font_family": "Arial", "palette": "business"}}
    personal = compile_design(base, snapshot)
    assert personal["typography"]["fontFamily"] == "Arial"
    assert base["typography"]["fontFamily"] == "Microsoft YaHei"
    protected = compile_design(base, snapshot, explicit_style=True)
    assert protected["palette"] == base["palette"]
    assert protected["typography"] == base["typography"]
