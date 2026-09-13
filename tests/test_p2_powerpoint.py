import io
import zipfile

from app.db.models import DeckSpecRecord, SlideSpecRecord
from app.db.session import SessionLocal
from app.powerpoint.roundtrip import apply_generated_effects, inspect_pptx
from app.templates.native_fill import fill_native_template
from pptx import Presentation


def _seed_slide(project_id: str) -> str:
    slide_id = "slide-p2"
    spec = {
        "id": slide_id, "projectId": project_id, "position": 1, "role": "content",
        "message": "高忠实往返保证既有品牌资产不丢失",
        "content": {"title": "PowerPoint 平台深度", "bullets": ["母版保留", "对象续编", "固定评测"]},
        "sourceRefs": [], "assetBindings": [],
        "visualIntent": {"primaryVisual": "typography"}, "designSystem": {},
        "speakerIntent": {"talkingPoints": ["说明高忠实往返"], "transition": "进入评测"},
    }
    with SessionLocal() as db:
        db.add(DeckSpecRecord(project_id=project_id, narrative={}, design_system={}, reproducibility={}))
        db.add(SlideSpecRecord(id=slide_id, project_id=project_id, position=1, spec=spec))
        db.commit()
    return slide_id


def _source_pptx(tmp_path) -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "既有企业演示"
    slide.placeholders[1].text = "原始正文"
    source = tmp_path / "source.pptx"
    deck.save(source)
    apply_generated_effects(source, [{"powerPoint": {"transition": "fade", "animation": "fade"}}])
    return source.read_bytes()


def test_generated_powerpoint_effects_and_narration_are_exported(client):
    project = client.post("/api/v1/projects", json={"name": "P2 effects"}).json()
    slide_id = _seed_slide(project["id"])
    effects = client.post(
        f"/api/v1/projects/{project['id']}/slides/{slide_id}/powerpoint-effects",
        json={"transition": "fade", "transition_speed": "med", "animation": "fade"},
    )
    assert effects.status_code == 200
    narration = client.post(
        f"/api/v1/projects/{project['id']}/narration",
        json={"slide_id": slide_id, "locale": "zh-CN", "words_per_minute": 220},
    )
    assert narration.status_code == 200
    assert narration.json()["slides"][0]["narration"].startswith("这一页说明")
    exported = client.get(f"/api/v1/projects/{project['id']}/export/pptx")
    assert exported.status_code == 200
    analysis = inspect_pptx(exported.content)
    assert analysis["transitionSlides"] == 1
    assert analysis["animationSlides"] == 1
    with zipfile.ZipFile(io.BytesIO(exported.content)) as package:
        notes = package.read("ppt/notesSlides/notesSlide1.xml").decode("utf-8")
    assert "旁白：" in notes and "[Sources]" in notes


def test_imported_pptx_object_edit_preserves_master_theme_and_effects(client, tmp_path):
    project = client.post("/api/v1/projects", json={"name": "P2 roundtrip"}).json()
    _seed_slide(project["id"])
    source = _source_pptx(tmp_path)
    imported = client.post(
        f"/api/v1/projects/{project['id']}/powerpoint/import",
        files={"file": ("company-deck.pptx", source, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert imported.status_code == 201
    original = imported.json()["analysis"]
    title = next(item for item in original["slides"][0]["objects"] if item["text"] == "既有企业演示")
    edited = client.post(
        f"/api/v1/projects/{project['id']}/powerpoint/slides/1/objects/{title['id']}",
        json={"value": "既有企业演示｜对象级续编"},
    )
    assert edited.status_code == 200
    assert edited.json()["slide"]["title"] == "既有企业演示｜对象级续编"
    exported = client.get(f"/api/v1/projects/{project['id']}/powerpoint/export")
    assert exported.status_code == 200
    result = inspect_pptx(exported.content)
    assert result["fidelityHashes"] == original["fidelityHashes"]
    assert result["transitionSlides"] == original["transitionSlides"] == 1
    assert result["animationSlides"] == original["animationSlides"] == 1
    template_output = tmp_path / "template-filled.pptx"
    template_report = fill_native_template(source, [{
        "position": 1, "role": "content", "message": "模板往返",
        "content": {"title": "模板往返", "bullets": ["保留主题和母版"]},
        "assetBindings": [],
    }], template_output)
    template_analysis = inspect_pptx(template_output.read_bytes())
    assert template_analysis["fidelityHashes"] == original["fidelityHashes"]
    assert template_report["themeMasterRoundtrip"] == "byte-stable"


def test_fixed_suite_and_blind_review_workflow(client, tmp_path):
    project = client.post("/api/v1/projects", json={"name": "P2 evaluation"}).json()
    _seed_slide(project["id"])
    assert client.get(f"/api/v1/projects/{project['id']}/export/pptx").status_code == 200
    source = _source_pptx(tmp_path)
    assert client.post(
        f"/api/v1/projects/{project['id']}/powerpoint/import",
        files={"file": ("comparison.pptx", source, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    ).status_code == 201
    run = client.post(f"/api/v1/projects/{project['id']}/evaluations")
    assert run.status_code == 201
    payload = run.json()
    assert payload["metrics"]["fixedCasesTotal"] == 11
    assert payload["suite"] == "professional-fixed-v2"
    assert payload["metrics"]["blindReady"] is True
    token = payload["blindToken"]
    blind = client.get(f"/api/v1/evaluations/blind/{token}")
    assert blind.status_code == 200 and blind.json()["candidateLabels"] == ["A", "B"]
    assert client.get(f"/api/v1/evaluations/blind/{token}/candidate/A").status_code == 200
    review = client.post(
        f"/api/v1/evaluations/blind/{token}/reviews",
        json={
            "reviewer_alias": "评审甲", "candidate_label": "A",
            "scores": {"fundamentals": 88, "visualDesign": 86, "completeness": 84, "correctness": 92, "fidelity": 90},
            "comment": "整体可交付",
        },
    )
    assert review.status_code == 201
    assert review.json()["average"]["fidelity"] == 90
