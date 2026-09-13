import io

from app.templates.analyzer import analyze_reference_pptx
from pptx import Presentation


def reference_deck() -> bytes:
    deck = Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = "公司战略沟通"
    cover.placeholders[1].text = "2026 年度规划"
    data = deck.slides.add_slide(deck.slide_layouts[5])
    data.shapes.title.text = "实验结果与核心指标"
    table = data.shapes.add_table(3, 3, 800000, 1700000, 7000000, 2600000).table
    for row, values in enumerate(
        (("方法", "Accuracy", "F1"), ("基线", "0.81", "0.80"), ("方法", "0.90", "0.93"))
    ):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    conclusion = deck.slides.add_slide(deck.slide_layouts[1])
    conclusion.shapes.title.text = "结论与下一步建议"
    conclusion.placeholders[1].text = "扩大验证范围\n完善证据链"
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def test_reference_analyzer_learns_page_functions_and_rhythm():
    analysis = analyze_reference_pptx(reference_deck())
    assert analysis["version"] == "reference-analysis-v3"
    assert analysis["constraintGrammar"]["version"] == "template-constraint-grammar-v1"
    assert "safeMargins" in analysis["constraintGrammar"]
    assert [slide["function"] for slide in analysis["slides"]] == [
        "cover",
        "data",
        "conclusion",
    ]
    assert analysis["slides"][1]["silhouette"]["hasTable"] is True
    assert analysis["rhythm"]["sequence"] == ["cover", "data", "conclusion"]
    assert analysis["rhythm"]["transitions"][0][0] == "cover→data"
    assert analysis["variantByRole"]["data"][1] == "table-highlight"
