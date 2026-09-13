import io
import json
from datetime import datetime

import fitz
from app.api.workflow_routes import clean_audience_bullet, compact_strategy_material
from app.documents import parse_source
from app.intelligence.content_selection import retain_source_boundaries, select_key_points
from app.intelligence.evidence import extract_evidence_graph
from app.intelligence.retrieval import retrieve_source_context
from app.presentation_intelligence.assets import _figure_ref
from app.presentation_intelligence.planner import plan_deck
from app.presentation_intelligence.small_model import bounded_batch_context
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


def saved(document):
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_word_preserves_first_heading_ancestry_table_order_and_boundaries(tmp_path):
    doc = Document()
    doc.add_heading("试点报告", 1)
    doc.add_heading("验证结果", 2)
    doc.add_paragraph("试点采用同一批样本复核。")
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ["方案", "准确率（%）", "耗时（秒）"]):
        cell.text = text
    for values in [["基线", "71.2", "12"], ["新方案", "89.6", ""]]:
        for cell, text in zip(table.add_row().cells, values):
            cell.text = text
    doc.add_paragraph("结果表明新方案提高了准确率，但仅限内部样本，不代表外部场景效果。")
    image = io.BytesIO()
    Image.new("RGB", (160, 100), "white").save(image, "PNG")
    doc.add_picture(io.BytesIO(image.getvalue()))
    doc.add_paragraph("图 1 试点框架")
    source = parse_source("试点.docx", saved(doc), asset_dir=tmp_path)
    section = source["sections"][1]
    assert source["sections"][0]["title"] == "试点报告"
    assert section["headingPath"] == ["试点报告", "验证结果"]
    assert [block["type"] for block in section["blocks"]][:3] == ["paragraph", "table", "paragraph"]
    assert "基线：准确率：71.2%；耗时：12秒" in section["text"]
    assert "新方案：准确率：89.6%；耗时：未提供" in section["text"]
    assert section["understanding"]["limitations"]
    assert source["tables"][0]["sourceRef"]["section"] == section["id"]
    assert source["figures"][0]["caption"] == "图 1 试点框架"
    assert _figure_ref(source, source["figures"][0])["section"] == section["id"]
    assert any("未提供" in fact["claim"] for fact in extract_evidence_graph([source])["facts"])


def test_spreadsheet_retains_tail_rows_percent_dates_and_missing_values():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["项目", "准确率", "日期", "复核结果"])
    for index in range(160):
        sheet.append([f"项目{index}", 0.896, datetime(2026, 9, 5), None])  # noqa: DTZ001 - Excel stores local dates.
        sheet.cell(index + 2, 2).number_format = "0.0%"
    source = parse_source("数据.xlsx", saved(workbook))
    text = "\n".join(section["text"] for section in source["sections"])
    assert "项目159" in text and "89.6%" in text and "2026-09-05" in text
    assert "复核结果：未提供" in text
    assert sum(table["rowCount"] for table in source["tables"]) == 160
    assert not source["readingWarnings"]
    json.dumps(source, ensure_ascii=False)


def test_markdown_table_and_chinese_csv_keep_cell_meaning():
    markdown = "# 报告\n## 结果\n| 方案 | 准确率（%） |\n| --- | ---: |\n| 基线 | 71.2 |\n| 新方案 | 89.6 |"
    source = parse_source("report.md", markdown.encode())
    assert source["sections"][1]["headingPath"] == ["报告", "结果"]
    assert "新方案：准确率：89.6%" in source["sections"][1]["text"]
    assert source["tables"][0]["headers"] == ["方案", "准确率（%）"]
    csv_source = parse_source("数据.csv", "方案,准确率\n基线,71.2%\n新方案,89.6%".encode("gb18030"))
    assert "新方案：准确率：89.6%" in csv_source["sections"][0]["text"]


def test_powerpoint_tables_are_read_into_slide_evidence():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "验证结果"
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(7), Inches(2)).table
    for row, values in zip(table.rows, [["方案", "准确率"], ["基线", "71.2%"], ["新方案", "89.6%"]]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    source = parse_source("旧稿.pptx", saved(presentation))
    assert source["sections"][0]["title"] == "验证结果"
    assert "新方案：准确率：89.6%" in source["sections"][0]["text"]
    assert source["tables"][0]["sourceRef"]["page"] == 1


def test_pdf_detects_cells_and_reports_unreadable_pages():
    document = fitz.open()
    page = document.new_page()
    page.insert_text((60, 60), "Evaluation results", fontsize=20)
    for x in [60, 200, 340]:
        page.draw_line((x, 100), (x, 220))
    for y in [100, 140, 180, 220]:
        page.draw_line((60, y), (340, y))
    for index, values in enumerate([["Method", "Accuracy"], ["Baseline", "71.2%"], ["New", "89.6%"]]):
        for column, value in enumerate(values):
            page.insert_text((70 + column * 140, 125 + index * 40), value)
    document.new_page()
    source = parse_source("results.pdf", document.tobytes())
    assert source["tables"][0]["rows"][2] == ["New", "89.6%"]
    assert any("New：Accuracy：89.6%" in section["text"] for section in source["sections"])
    assert any(warning["code"] == "ocr-required" and warning["page"] == 2 for warning in source["readingWarnings"])


def long_section(section_id="S001"):
    return {"id": section_id, "title": "试点验证", "text": "介绍这次材料整理的工作背景。" * 90
            + "结果表明新方案减少了人工复核。仅限内部样本，不能推断外部效果。"}


def test_extractive_selection_keeps_tail_findings_and_qualitative_boundaries():
    section = long_section()
    points = select_key_points(section, maximum=3)
    assert "结果表明新方案减少了人工复核。" in points
    assert "仅限内部样本，不能推断外部效果。" in points
    graph = extract_evidence_graph([{"name": "brief.md", "sections": [section]}])
    assert any(fact["kind"] == "constraint" for fact in graph["facts"])
    claim = "结果显示新方案提高效率，" + "核对过程保留人工确认，" * 8 + "但仅限内部试点，尚未验证外部场景。"
    assert "尚未验证外部场景" in clean_audience_bullet(claim)
    signed = extract_evidence_graph([{"name": "data", "sections": [{"id": "S1", "text": "净增长为 -12.5%。"}]}])
    assert signed["metrics"][0]["value"] == -12.5
    protected = retain_source_boundaries(["仅限内部样本，不能推断外部效果。"], ["方案提高效率", "推广到各个团队", "降低复核成本", "加快业务响应"], 4)
    assert len(protected) == 4 and protected[-1] == "仅限内部样本，不能推断外部效果。"


def test_bounded_context_covers_every_fixed_reference_and_excludes_unrelated_material():
    source = {"name": "source.md", "sections": [long_section(f"S{index:03d}") for index in range(1, 4)]}
    unrelated = {"name": "unrelated.md", "sections": [{"id": "S001", "text": "无关材料不能填充到这批页面。"}]}
    slides = [{"message": "试点结果", "sourceRefs": [{"document": source["name"], "section": section["id"]}]} for section in source["sections"]]
    context = bounded_batch_context([source, unrelated], slides, "试点", "business", "", 700)
    assert len(context) <= 700
    for section in source["sections"]:
        assert section["id"] in context
    assert context.count("不能推断外部效果") == 3
    assert "unrelated" not in context


def test_strategy_context_is_valid_json_and_includes_late_conclusions():
    source = {"name": "报告.md", "sections": [{"id": f"S{index:03d}", "title": "背景材料", "text": "普通背景说明。" * 40} for index in range(30)]}
    source["sections"].append({"id": "S031", "title": "结论与边界", "text": "建议先在内部推广。外部效果尚未验证。"})
    context = compact_strategy_material({}, [source], 1000)
    assert len(context) <= 1000
    assert json.loads(context)["omittedSections"] > 0
    assert "外部效果尚未验证" in context
    retrieved = retrieve_source_context([source], "结论", "business", "", 800)
    assert len(retrieved) <= 800 and "S031" in retrieved


def test_short_business_deck_uses_data_and_keeps_plain_file_content():
    source = parse_source("材料.md", "# 背景\n业务需要减少人工复核的等待。\n# 方法\n首先整理资料。然后核对来源。最后由专人确认。\n# 验证结果\n基线：准确率：71.2%。新方案：准确率：89.6%。\n# 结论\n建议先在内部推广，但不能外推到其他行业。".encode())
    plan = plan_deck([source], "试点汇报", "business", 5)
    assert any(slide["role"] == "data" for slide in plan["slides"])
    assert "不能外推" in " ".join(plan["slides"][-1]["content"]["bullets"])
    plain = parse_source("notes.txt", "结果表明材料理解是制作演示的基础。".encode())
    plain_plan = plan_deck([plain], "普通笔记", "business", 3)
    assert plain_plan["slides"][1]["sourceRefs"]


def test_upload_and_workspace_expose_reading_diagnostics(client):
    project = client.post("/api/v1/projects", json={"name": "读取检查"}).json()
    document = fitz.open()
    document.new_page()
    uploaded = client.post(f"/api/v1/projects/{project['id']}/sources", files={"file": ("扫描.pdf", document.tobytes(), "application/pdf")})
    assert uploaded.status_code == 201
    assert uploaded.json()["readingStatus"] == "needs-attention"
    workspace = client.get(f"/api/v1/projects/{project['id']}/workspace").json()
    assert workspace["sources"][0]["readingWarnings"][0]["code"] == "ocr-required"
    outline = client.post(f"/api/v1/projects/{project['id']}/outline", json={"title": "扫描报告", "preset": "business", "slide_count": 6})
    assert outline.status_code == 422 and "没有可读取的正文" in outline.json()["detail"]


def test_missing_formula_cache_is_unknown_and_literal_percent_is_not_scaled():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["方案", "实际值", "公式结果"])
    sheet.append(["方案甲", 12, "=1+2"])
    sheet.cell(2, 2).number_format = '0"%"'
    source = parse_source("公式.xlsx", saved(workbook))
    assert "实际值：12" in source["sections"][0]["text"]
    assert "1200%" not in source["sections"][0]["text"]
    assert "公式结果：未提供" in source["sections"][0]["text"]
    assert source["readingWarnings"][0]["code"] == "formula-result-unavailable"


def test_long_table_limit_is_reported_and_empty_cells_are_not_zero():
    csv = "对象,数值\n" + "\n".join(f"项目{index}," for index in range(10010))
    source = parse_source("large.csv", csv.encode())
    assert source["readingWarnings"][0]["code"] == "row-limit"
    assert sum(table["rowCount"] for table in source["tables"]) == 9999
    assert "数值：未提供" in source["sections"][-1]["text"]
