import asyncio
import copy
import importlib.util
import json
import shutil
from pathlib import Path

import pytest
from app.db.models import DeckSpecRecord, SlideSpecRecord
from app.db.session import SessionLocal
from app.documents import parse_source
from app.evaluation.common import parse_json_object, slide_text
from app.evaluation.ppteval import evaluate_ppteval
from app.evaluation.quiz import evaluate_quiz, lexical_answer, number_keys, rule_questions
from app.evaluation.render_review import analyze_page_image, review_rendered_pptx
from app.presentation_intelligence.planner import plan_deck
from app.validation.quality import validate_deck
from PIL import Image, ImageDraw
from sqlalchemy import select

ROOT = Path(__file__).resolve().parents[1]


def _fixture(name="data-dashboard", preset="data"):
    path = ROOT / "benchmarks" / name / "fixture.md"
    source = parse_source(path.name, path.read_bytes(), "text/markdown")
    plan = plan_deck([source], name, preset, 8)
    return source, plan["slides"]


class FakeJudge:
    def __init__(self, replies=None, vision=False, fail=False):
        self.model_id = "fake-judge"
        self.vision = vision
        self.replies = list(replies or [])
        self.fail = fail
        self.prompts = []
        self.image_batches = []

    async def complete(self, prompt, max_tokens=2400):
        self.prompts.append(prompt)
        if self.fail:
            raise RuntimeError("provider down")
        return self.replies.pop(0) if self.replies else "{}"

    async def see(self, prompt, images, max_tokens=2400):
        self.image_batches.append(len(images))
        return await self.complete(prompt, max_tokens)


def test_parse_json_object_repairs_common_model_damage():
    assert parse_json_object('<think>x</think>```json\n{"a": 1,}\n```') == {"a": 1}
    assert parse_json_object('结果如下：{"slides":[{"position":1,"score":4},]} 谢谢') == {
        "slides": [{"position": 1, "score": 4}]
    }
    assert parse_json_object('{"score": 4, "rationale": “清晰”}') == {"score": 4, "rationale": "清晰"}
    truncated = parse_json_object('{"slides":[{"position":2,"rationale":"内容完')
    assert truncated["slides"][0]["position"] == 2
    with pytest.raises(ValueError):
        parse_json_object("没有任何结构")


def test_ppteval_rule_mode_scores_every_dimension_with_chinese_rationale():
    source, slides = _fixture()
    report = validate_deck(slides, [source])
    result = asyncio.run(evaluate_ppteval(slides, report))
    assert result["mode"] == "rule" and result["scale"] == [1, 5]
    assert len(result["slides"]) == len(slides)
    for dimension in ("content", "design", "coherence"):
        assert 1 <= result["deck"][dimension]["score"] <= 5
        assert result["deck"][dimension]["rationale"]
        for item in result["slides"]:
            assert item[dimension]["score"] in {1, 2, 3, 4, 5}
            assert item[dimension]["source"] == "rule"

    dense = copy.deepcopy(slides)
    dense[3]["content"]["bullets"] = ["这是一条很长的观众文案" * 12] * 6
    dense_report = validate_deck(dense, [source])
    worse = asyncio.run(evaluate_ppteval(dense, dense_report))
    page = next(item for item in worse["slides"] if item["position"] == 4)
    assert page["design"]["score"] < 4
    assert any(issue["dimension"] == "design" and issue["suggestion"] for issue in page["issues"])


def test_ppteval_render_issues_lower_design_on_real_pages():
    source, slides = _fixture()
    report = validate_deck(slides, [source])
    render_pages = {2: {"score": 2, "issues": [{"code": "low-contrast", "severity": "error"}]}}
    result = asyncio.run(evaluate_ppteval(slides, report, render_pages=render_pages))
    page = next(item for item in result["slides"] if item["position"] == 2)
    assert page["design"]["score"] <= 2
    assert any(issue["code"] == "low-contrast" for issue in page["issues"])


def test_ppteval_text_judge_merges_valid_scores_and_keeps_rule_design():
    source, slides = _fixture()
    report = validate_deck(slides, [source])
    first_batch = json.dumps({"slides": [
        {"position": 1, "scores": {"content": {"score": 2, "rationale": "封面缺少主题信息", "suggestion": "补充副标题"},
                                   "coherence": {"score": 4, "rationale": "开场清晰"}}},
        {"position": 99, "scores": {"content": {"score": 5}}},
        {"position": 3, "scores": {"content": {"score": "非法"}}},
    ]}, ensure_ascii=False)
    judge = FakeJudge([
        "```json\n" + first_batch + "\n```",
        "not json at all",
        '{"score": 3, "rationale": "结尾行动建议较弱", "suggestions": ["第 8 页：补充下一步责任人"]}',
    ])
    result = asyncio.run(evaluate_ppteval(slides, report, text_judge=judge))
    assert result["mode"] == "text"
    cover = result["slides"][0]
    assert cover["content"] == {"score": 2, "rationale": "封面缺少主题信息", "source": "text"}
    assert cover["design"]["source"] == "rule"
    assert any(issue["suggestion"] == "补充副标题" for issue in cover["issues"])
    assert result["slides"][2]["content"]["source"] == "rule"
    assert result["deck"]["coherence"]["suggestions"] == ["第 8 页：补充下一步责任人"]
    assert any("文本评审失败" in error for error in result["errors"])

    failing = asyncio.run(evaluate_ppteval(slides, report, text_judge=FakeJudge(fail=True)))
    assert failing["mode"] == "rule" and failing["errors"]


def test_ppteval_vision_judge_reads_rendered_pptx_pages(tmp_path):
    source, slides = _fixture()
    report = validate_deck(slides, [source])
    render_pages = {}
    for slide in slides:
        path = tmp_path / f"page-{slide['position']}.png"
        Image.new("RGB", (64, 36), "white").save(path)
        render_pages[slide["position"]] = {"imagePath": str(path), "issues": [], "score": 5}
    replies = [
        json.dumps({"slides": [
            {"position": position, "scores": {name: {"score": 3, "rationale": "层级尚可"} for name in
                                              ("content", "design", "coherence")}}
            for position in range(start, start + 4)
        ]}, ensure_ascii=False)
        for start in (1, 5)
    ]
    vision = FakeJudge(replies + ['{"score": 4, "rationale": "逻辑递进"}'], vision=True)
    result = asyncio.run(evaluate_ppteval(slides, report, render_pages=render_pages, vision_judge=vision))
    assert vision.image_batches == [4, 4]
    assert result["mode"] == "vision" and result["model"] == "fake-judge"
    assert all(item["design"]["source"] == "vision" for item in result["slides"])
    assert result["deck"]["design"]["score"] == 3.0


def test_quiz_rule_questions_trace_and_missing_facts():
    source, slides = _fixture()
    questions = rule_questions([source], 6, 4)
    assert {item["type"] for item in questions} == {"detail", "concept"}
    assert all("____" in item["question"] for item in questions if item["type"] == "detail")
    assert number_keys("18,400 sessions") == number_keys("18400")
    assert number_keys("100 synthetic records") == {(100.0, "")}

    full = asyncio.run(evaluate_quiz(slides, [source]))
    assert full["status"] == "completed" and full["generator"] == "rule" and full["answerer"] == "lexical"
    assert full["retention"] == 100.0 and full["missingFacts"] == []

    trimmed = [slide for slide in slides if "Acquisition" not in str(slide["content"].get("title"))]
    partial = asyncio.run(evaluate_quiz(trimmed, [source]))
    assert partial["retention"] < full["retention"]
    missed = [item for item in partial["questions"] if not item["correct"]]
    assert missed and all(item["supportedBySlide"] is None for item in missed)
    assert all({"question", "expected", "answer", "supportedBySlide", "type"} <= set(item) for item in partial["questions"])
    assert any("16%" in fact["missingValues"] for fact in partial["missingFacts"])
    hit = next(item for item in partial["questions"] if item["correct"])
    assert hit["supportedBySlide"] in {slide["position"] for slide in trimmed}


def test_quiz_lexical_answer_rejects_wrong_number():
    slides = [{"position": 1, "content": {"title": "转化", "bullets": ["合作渠道转化率为 15%"]}}]
    question = {"type": "detail", "question": "合作渠道转化率是多少？", "expected": "16%",
                "context": "合作渠道转化率为"}
    answer = lexical_answer(question, slides)
    assert answer["answer"] == "15%" and answer["supportedBySlide"] is None


def test_quiz_llm_mode_requires_cited_slide_support():
    source, slides = _fixture()
    questions = {"questions": [
        {"type": "detail", "question": "整体转化率是多少？", "answer": "12%", "document": "fixture.md",
         "section": "S002", "importance": 3},
        {"type": "detail", "question": "满意度是多少？", "answer": "91%", "document": "fixture.md",
         "section": "S003", "importance": 2},
        {"type": "concept", "question": "数据有什么限制？", "answer": "The sample is synthetic and observational",
         "document": "fixture.md", "section": "S999"},
        {"type": "bogus", "question": "忽略", "answer": "x"},
    ]}
    overview = next(slide["position"] for slide in slides if "Conversion was 12%" in slide_text(slide))
    answers = {"answers": [
        {"id": "q1", "answer": "12%", "slide": overview},
        {"id": "q2", "answer": "91%", "slide": 1},  # cover does not contain 91%: unsupported
        {"id": "q3", "answer": "无法从页面回答", "slide": None},
    ]}
    judge = FakeJudge([json.dumps(questions, ensure_ascii=False), json.dumps(answers, ensure_ascii=False)])
    result = asyncio.run(evaluate_quiz(slides, [source], judge=judge))
    assert result["generator"] == "llm" and result["answerer"] == "llm"
    trace = {item["id"]: item for item in result["questions"]}
    assert trace["q1"]["correct"] and trace["q1"]["supportedBySlide"] == overview
    assert not trace["q2"]["correct"] and trace["q2"]["citedSlide"] == 1
    assert not trace["q3"]["correct"]
    assert trace["q3"]["sourceRef"]["section"] == "S001"
    assert result["detailRetention"] == 50.0 and result["conceptRetention"] == 0.0
    # The answering model sees slide text and questions, never the reference answers.
    assert "标准答案" not in judge.prompts[1] and "12%" in judge.prompts[1]

    brief_only = asyncio.run(evaluate_quiz(slides, [{"name": "brief", "synthetic": True, "sections": []}]))
    assert brief_only["status"] == "no-sources"


def _page(path: Path, draw=None, background="white"):
    image = Image.new("RGB", (960, 540), background)
    if draw:
        draw(ImageDraw.Draw(image))
    image.save(path)


def _fake_renderer(pages):
    def render(pptx, output):
        for index, painter in enumerate(pages, 1):
            _page(output / f"page-{index}.png", painter)
        return {"status": "rendered-needs-review", "fileHash": "abc", "pages": [
            {"page": index, "image": f"page-{index}.png", "text": text}
            for index, text in enumerate(["标题一 正文", "", "标题三", "标题四", "标题五"][: len(pages)], 1)
        ]}
    return render


def test_render_review_deterministic_checks_on_page_images(tmp_path):
    good = lambda draw: (draw.rectangle((80, 60, 700, 120), fill="black"),
                         draw.rectangle((80, 200, 860, 420), fill="#1f4e79"))
    def faint(draw):  # thin, nearly invisible strokes: text the audience cannot read
        for x in range(80, 700, 9):
            draw.rectangle((x, 60, x + 2, 140), fill="#f2f2f2")

    def panel(draw):  # a pale solid panel behind dark text is legitimate design
        draw.rectangle((60, 300, 900, 420), fill="#e9edf3")
        draw.rectangle((90, 340, 600, 370), fill="#1b2433")
    edge = lambda draw: (draw.rectangle((80, 60, 700, 120), fill="black"),
                         draw.rectangle((700, 240, 960, 300), fill="black"))
    pptx = tmp_path / "deck.pptx"
    pptx.write_bytes(b"pptx")
    slides = [{"position": index, "content": {"title": f"标题{name}"}} for index, name in enumerate("一二三四五", 1)]
    result = asyncio.run(review_rendered_pptx(
        pptx, slides, tmp_path / "render", renderer=_fake_renderer([good, None, faint, edge, panel]),
    ))
    assert result["status"] == "completed" and result["mode"] == "deterministic"
    codes = {page["page"]: {issue["code"] for issue in page["issues"]} for page in result["pages"]}
    assert "blank-page" not in codes[1] and "low-contrast" not in codes[1]
    assert "blank-page" in codes[2]
    assert "low-contrast" in codes[3]
    assert "edge-touching" in codes[4]
    assert codes[5] == set()
    assert result["errorCount"] >= 2
    assert analyze_page_image(tmp_path / "render" / "page-1.png")["metrics"]["contrast"] > 7


def test_render_review_vision_merge_and_graceful_fallbacks(tmp_path):
    pptx = tmp_path / "deck.pptx"
    pptx.write_bytes(b"pptx")
    slides = [{"position": 1, "content": {"title": "标题一"}}]
    painter = lambda draw: draw.rectangle((80, 60, 700, 300), fill="black")
    vision = FakeJudge([json.dumps({"pages": [{"page": 1, "score": 3, "issues": [
        {"type": "overlap", "severity": "error", "message": "图片压住标题", "suggestion": "下移图片"}]}]},
        ensure_ascii=False)], vision=True)
    result = asyncio.run(review_rendered_pptx(
        pptx, slides, tmp_path / "render", vision_judge=vision, renderer=_fake_renderer([painter]),
    ))
    assert result["mode"] == "vision" and vision.image_batches == [1]
    assert any(issue["code"] == "vision-overlap" and issue["suggestion"] == "下移图片" for issue in result["issues"])

    unavailable = asyncio.run(review_rendered_pptx(
        pptx, slides, tmp_path / "missing",
        renderer=lambda *_: {"status": "unavailable", "message": "本机未找到 LibreOffice"},
    ))
    assert unavailable["status"] == "unavailable" and unavailable["message"] == "本机未找到 LibreOffice"
    assert asyncio.run(review_rendered_pptx(None, slides, tmp_path / "none"))["status"] == "no-export"


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice missing")
def test_render_review_renders_real_pptx_with_libreoffice(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[6])
    box = first.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1.5))
    box.text_frame.text = "Quarterly results"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    deck.slides.add_slide(deck.slide_layouts[6])
    path = tmp_path / "real.pptx"
    deck.save(path)
    slides = [{"position": 1, "content": {"title": "Quarterly results"}}, {"position": 2, "content": {"title": ""}}]
    result = asyncio.run(review_rendered_pptx(path, slides, tmp_path / "render"))
    assert result["status"] == "completed" and result["pageCount"] == 2
    assert {issue["code"] for issue in result["pages"][1]["issues"]} == {"blank-page"}
    assert "blank-page" not in {issue["code"] for issue in result["pages"][0]["issues"]}


def _seed_project(client):
    project = client.post("/api/v1/projects", json={"name": "研究评测"}).json()
    fixture = ROOT / "benchmarks" / "data-dashboard" / "fixture.md"
    uploaded = client.post(
        f"/api/v1/projects/{project['id']}/sources",
        files={"file": ("fixture.md", fixture.read_bytes(), "text/markdown")},
    )
    assert uploaded.status_code == 201
    source = parse_source(fixture.name, fixture.read_bytes(), "text/markdown")
    slides = plan_deck([source], "data-dashboard", "data", 8)["slides"]
    with SessionLocal() as db:
        db.add(DeckSpecRecord(project_id=project["id"], narrative={}, design_system={}, reproducibility={}))
        for slide in slides:
            db.add(SlideSpecRecord(id=slide["id"], project_id=project["id"], position=slide["position"],
                                   spec={**slide, "projectId": project["id"]}))
        db.commit()
    return project["id"], slides


def test_evaluation_api_persists_runs_and_flags_stale_results(client, monkeypatch):
    project_id, _ = _seed_project(client)
    response = client.post(f"/api/v1/projects/{project_id}/evaluations",
                           json={"kinds": ["ppteval", "quiz", "render"], "use_models": False})
    assert response.status_code == 201
    run = response.json()
    assert run["status"] == "completed" and run["suite"] == "research-v1" and run["stale"] is False
    assert run["kinds"] == ["ppteval", "quiz", "render"]
    assert run["ppteval"]["mode"] == "rule" and 1 <= run["summary"]["content"] <= 5
    assert run["quiz"]["retention"] == run["summary"]["retention"] and run["quiz"]["questions"]
    assert run["render"]["status"] == "no-export"

    listed = client.get(f"/api/v1/projects/{project_id}/evaluations").json()
    assert [item["id"] for item in listed] == [run["id"]] and "ppteval" not in listed[0]
    workspace = client.get(f"/api/v1/projects/{project_id}/workspace").json()
    assert all(item["suite"] != "research-v1" for item in workspace["evaluationRuns"])

    with SessionLocal() as db:
        row = db.scalar(select(SlideSpecRecord).where(SlideSpecRecord.project_id == project_id,
                                                       SlideSpecRecord.position == 4))
        row.spec = {**row.spec, "content": {**row.spec["content"], "title": "修改后的标题"}}
        db.commit()
    detail = client.get(f"/api/v1/projects/{project_id}/evaluations/{run['id']}").json()
    assert detail["stale"] is True and detail["quiz"]["questions"]
    assert client.get(f"/api/v1/projects/{project_id}/evaluations/missing").status_code == 404
    assert client.get(f"/api/v1/projects/{project_id}/evaluations/{run['id']}/pages/1").status_code == 404

    from app.api import quality_routes

    judge = FakeJudge(fail=True)
    monkeypatch.setattr(quality_routes, "resolve_judges", lambda db: (judge, None))
    with_model = client.post(f"/api/v1/projects/{project_id}/evaluations", json={"kinds": ["quiz"]}).json()
    assert with_model["status"] == "completed" and with_model["models"]["text"] == "fake-judge"
    assert with_model["quiz"]["generator"] == "rule" and with_model["quiz"]["errors"]
    assert "ppteval" not in with_model

    rejected = client.post(f"/api/v1/projects/{project_id}/evaluations", json={"kinds": ["unknown"]})
    assert rejected.status_code == 422


def test_evaluate_deck_cli_is_deterministic(tmp_path, capsys):
    spec = importlib.util.spec_from_file_location("evaluate_deck", ROOT / "scripts" / "evaluate-deck.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    fixture = str(ROOT / "benchmarks" / "data-dashboard" / "fixture.md")
    output = tmp_path / "result.json"
    assert module.main([fixture, "--preset", "data", "--out", str(output)]) == 0
    summary = json.loads(capsys.readouterr().out)["summary"]
    assert summary["pptevalMode"] == "rule" and summary["retention"] == 100.0
    first = json.loads(output.read_text(encoding="utf-8"))
    deck = tmp_path / "deck.json"
    _, slides = _fixture()
    deck.write_text(json.dumps({"slides": slides}, ensure_ascii=False), encoding="utf-8")
    assert module.main([fixture, "--deck", str(deck), "--kinds", "quiz", "--out", str(output)]) == 0
    second = json.loads(output.read_text(encoding="utf-8"))
    assert second["kinds"] == ["quiz"]
    assert second["quiz"]["questions"] == first["quiz"]["questions"]
