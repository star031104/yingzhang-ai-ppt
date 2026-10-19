import io
import json
from types import SimpleNamespace

from app.presentation_intelligence.layout_planner import plan_deck_layouts
from app.templates.analyzer import analyze_reference_pptx
from app.templates.layout_library import (
    attach_reference_layouts,
    load_reference_layout_library,
    match_reference_layout,
)
from app.templates.native_fill import fill_native_template
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

PRIVATE_TEXT = ("年度经营回顾", "推进渠道下沉与重点客户共建", "季度收入表现", "总结与下一步")


def reference_deck() -> bytes:
    deck = Presentation()
    width, height = int(deck.slide_width), int(deck.slide_height)
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = PRIVATE_TEXT[0]
    cover.placeholders[1].text = "经营管理部"
    content = deck.slides.add_slide(deck.slide_layouts[4])  # "Comparison" master layout
    content.shapes.title.text = "三项举措"
    content.placeholders[1].text_frame.text = PRIVATE_TEXT[1] + "，形成区域示范并持续复制。"
    for line in ("建设统一数据平台，打通线索、合同与回款的全过程记录。",
                 "优化交付流程，把实施周期从季度级压缩到月度级并稳定运行。"):
        content.placeholders[1].text_frame.add_paragraph().text = line
    data = deck.slides.add_slide(deck.slide_layouts[5])
    data.shapes.title.text = PRIVATE_TEXT[2]
    table = data.shapes.add_table(3, 3, 800000, 1700000, 7000000, 2600000).table
    for row, values in enumerate((("季度", "收入", "增速"), ("Q1", "1200万", "12%"), ("Q2", "1500万", "25%"))):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    picture = deck.slides.add_slide(deck.slide_layouts[6])
    image = io.BytesIO()
    Image.new("RGB", (64, 36), (40, 90, 160)).save(image, format="PNG")
    image.seek(0)
    picture.shapes.add_picture(image, Emu(width // 10), Emu(height // 10), Emu(width * 8 // 10), Emu(height * 7 // 10))
    closing = deck.slides.add_slide(deck.slide_layouts[1])
    closing.shapes.title.text = PRIVATE_TEXT[3]
    closing.placeholders[1].text = "扩大试点范围\n建立季度复盘机制"
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def spec(position, role, bullets=(), **extra):
    return {"position": position, "role": role, "message": "", "purpose": role,
            "content": {"title": f"页面{position}", "bullets": list(bullets)},
            "assetBindings": [], "visualIntent": {"archetypeCandidates": ["split", "statement"]}, **extra}


def planned_slides(image_path=None):
    return [
        spec(1, "cover"),
        spec(2, "method", ["梳理客户分层与触达节奏", "统一线索与合同数据", "压缩交付周期并复盘"]),
        spec(3, "data", ["Q1 收入 1200万，增速 12%", "Q2 收入 1500万，增速 25%"],
             visualIntent={"archetypeCandidates": ["chart-focus", "metric-wall"], "expectedFigureKinds": ["chart", "table"]}),
        spec(4, "evidence", ["现场照片展示门店改造效果"],
             assetBindings=[{"type": "source-image", "path": str(image_path or "")}]),
        spec(5, "conclusion", ["扩大试点", "建立复盘机制"]),
    ]


def library():
    analysis = analyze_reference_pptx(reference_deck())
    tagged = [{**entry, "sourceKind": "skill", "sourceId": "ref", "sourceName": "参考"} for entry in analysis["layoutLibrary"]["entries"]]
    return analysis, {**analysis["layoutLibrary"], "entries": tagged}


def test_analyzer_persists_text_free_layout_signatures():
    analysis, _ = library()
    layout_library = analysis["layoutLibrary"]
    entries = layout_library["entries"]
    assert layout_library["version"] == "reference-layout-library-v1"
    assert [entry["function"] for entry in entries] == ["cover", "content", "data", "image-story", "conclusion"]
    content, data, picture = entries[1], entries[2], entries[3]
    assert content["layout"] == "Comparison" and content["layoutIndex"] == 4 and content["bulletCount"] == 3
    assert data["tableCount"] == 1 and data["metricCount"] >= 4 and data["variantHint"] == "table-highlight"
    assert picture["pictureCount"] == 1 and picture["silhouette"]["mediaAreaRatio"] > 0.3
    assert {region["kind"] for region in data["regions"]} >= {"title", "table"}
    assert all(0 <= region["x"] <= 1 and 0 < region["w"] <= 1 for entry in entries for region in entry["regions"])
    serialized = json.dumps(layout_library, ensure_ascii=False)
    assert not any(text in serialized for text in PRIVATE_TEXT)


def test_pages_match_reference_slides_by_content_signature(tmp_path):
    _, tagged = library()
    image = tmp_path / "photo.png"
    Image.new("RGB", (8, 8)).save(image)
    slides = planned_slides(image)
    summary = attach_reference_layouts(slides, tagged)
    assert summary["matched"] == 5
    assert [slide["referenceLayout"]["slidePosition"] for slide in slides] == [1, 2, 3, 4, 5]
    assert slides[2]["referenceLayout"]["variantHint"] == "table-highlight"
    assert "参考页含图表或表格" in slides[2]["referenceLayout"]["reasons"]
    # A metric-free page never borrows a table layout; an unrelated role finds nothing.
    assert match_reference_layout(spec(9, "agenda", ["一", "二"]), tagged["entries"]) is None
    layout = plan_deck_layouts(slides, "business")
    data_plan = slides[2]["layoutPlan"]
    assert data_plan["recommendedVariant"] == "table-highlight"
    assert data_plan["referenceLayout"]["followed"] is True
    assert data_plan["constraints"]["referenceSafeMargins"]
    assert len(layout["familySequence"]) == 5


def test_native_fill_prefers_the_matched_master_layout(tmp_path):
    _, tagged = library()
    slides = planned_slides()
    attach_reference_layouts(slides, tagged)
    template = reference_deck()
    report = fill_native_template(template, slides, tmp_path / "matched.pptx")
    layouts = [slide.slide_layout.name for slide in Presentation(tmp_path / "matched.pptx").slides]
    assert layouts[1] == "Comparison"
    assert report["referenceLayoutMatches"] >= 3 and not report["unmatchedSlides"]
    plain = [{key: value for key, value in slide.items() if key != "referenceLayout"} for slide in slides]
    fill_native_template(template, plain, tmp_path / "plain.pptx")
    assert Presentation(tmp_path / "plain.pptx").slides[1].slide_layout.name != "Comparison"


def test_library_loader_merges_skill_and_personal_reference(client):
    from app.personalization.private_files import reference_root
    from test_personal_experience import fresh
    from test_personalization import make_profile, plan

    profile = make_profile(client)
    base = f"/api/v1/me/profiles/{profile['id']}"
    response = client.post(base + "/reference-template", data={"revision": profile["revision"], "retain_original": "true"},
                           files={"file": ("reference.pptx", reference_deck())})
    assert response.status_code == 201, response.text
    reference = response.json()
    assert reference["analysis"]["layoutLibrary"]["entries"]
    assert not any(text in response.text for text in PRIVATE_TEXT)
    profile = fresh(client, profile)
    client.post(base + f"/references/{reference['id']}/activate", json={"revision": profile["revision"]})
    profile = fresh(client, profile)
    _project, result = plan(client, profile)
    assert result["layoutPlanning"]["referenceLayouts"]["matched"] >= 2
    matched = [slide["referenceLayout"] for slide in result["slides"] if slide.get("referenceLayout")]
    assert {item["source"]["kind"] for item in matched} == {"personal-reference"}
    assert result["slides"][0]["referenceLayout"]["function"] == "cover"

    from app.db.models import PersonalReference
    from app.db.session import SessionLocal
    with SessionLocal() as db:
        row = db.get(PersonalReference, reference["id"])
        assert (reference_root(row) / "reference.pptx").is_file()
        row.analysis = {key: value for key, value in row.analysis.items() if key != "layoutLibrary"}
        db.commit()
        skill = SimpleNamespace(id="reference-demo", manifest={
            "name": "参考技能", "analysis": analyze_reference_pptx(reference_deck()),
        })
        loaded = load_reference_layout_library(db, [skill], {"reference": {"id": reference["id"]}})
    kinds = [entry["sourceKind"] for entry in loaded["entries"]]
    assert kinds.count("personal-reference") == 5 and kinds.count("skill") == 5
    assert [source["kind"] for source in loaded["sources"]] == ["personal-reference", "skill"]
    assert loaded["constraintGrammar"]["safeMargins"]
