import copy
import hashlib
import io
import json
import zipfile
from pathlib import Path

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from pptx import Presentation

from app.config import settings
from app.db.models import PersonalBinding, PersonalCase, PersonalComparison, PersonalReference, Project
from app.db.session import SessionLocal
from app.main import app
from app.personalization.private_files import validate_pptx
from app.personalization.roundtrip_learning import annotate_export, compare_external
from test_personalization import make_profile, teach, plan, reset


def pptx_bytes():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = 'Private original'
    slide.placeholders[1].text = 'Evidence stays private'
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def fresh(client, profile):
    return next(row for row in client.get('/api/v1/me/profiles').json()['profiles'] if row['id'] == profile['id'])


def test_retained_reference_consent_activation_and_reset(client):
    profile = make_profile(client)
    base = f"/api/v1/me/profiles/{profile['id']}"
    data = pptx_bytes()
    assert client.post(base + '/reference-template', data={'revision': profile['revision']}, files={'file': ('sample.pptx', data)}).status_code == 422
    response = client.post(base + '/reference-template', data={'revision': profile['revision'], 'retain_original': 'true'}, files={'file': ('sample.pptx', data)})
    assert response.status_code == 201, response.text
    ref = response.json()
    assert not ref['active'] and 'Evidence stays private' not in response.text
    profile = fresh(client, profile)
    assert client.post(base + f"/references/{ref['id']}/activate", json={'revision': profile['revision']}).status_code == 200
    profile = fresh(client, profile)
    project, _ = plan(client, profile)
    with SessionLocal() as db:
        assert db.get(PersonalBinding, project['id']).snapshot['reference']['id'] == ref['id']
        path = Path(db.get(PersonalReference, ref['id']).artifact_path)
        assert (path / 'reference.pptx').is_file()
    reset(client)
    assert not path.exists()
    with SessionLocal() as db:
        assert db.get(PersonalReference, ref['id']) is None
        assert db.get(Project, project['id']) is not None


def test_pack_import_is_paused_and_does_not_trust_quality(client):
    profile = make_profile(client)
    teach(client, profile, 'font_family', 'Arial')
    with SessionLocal() as db:
        from app.personalization.cases import extract_case
        features = extract_case([{'role': 'cover', 'content': {'title': 'Secret phrase'}}])
        features['quality'] = {'passed': True}
        db.add(PersonalCase(owner_id=db.get(__import__('app.db.models', fromlist=['PersonalProfile']).PersonalProfile, profile['id']).owner_id, profile_id=profile['id'], label='Example', scenario='business', features=features, status='confirmed', source_hash='example'))
        db.commit()
    response = client.post(f"/api/v1/me/profiles/{profile['id']}/experience-pack", json={})
    assert response.status_code == 200
    payload = response.json()
    assert 'Secret phrase' not in response.text
    result = client.post('/api/v1/me/experience-pack-import', files={'file': ('experience.json', json.dumps(payload).encode())})
    assert result.status_code == 201, result.text
    imported = result.json()
    assert not imported['use_memory']
    experience = client.get(f"/api/v1/me/profiles/{imported['id']}/experience").json()
    case = experience['cases'][0]
    assert case['status'] == 'candidate' and not case['features']['quality']['passed']
    endpoint = f"/api/v1/me/profiles/{imported['id']}/cases/{case['id']}/confirm"
    assert client.post(endpoint, json={'revision': imported['revision']}).status_code == 422
    assert client.post(endpoint, json={'revision': imported['revision'], 'reviewed_imported_example': True}).status_code == 200
    broken = copy.deepcopy(payload); broken['name'] = 'tampered'
    assert client.post('/api/v1/me/experience-pack-import', files={'file': ('bad.json', json.dumps(broken).encode())}).status_code == 422


def test_case_quality_gate_and_forgetting_derived_memory(client, monkeypatch):
    from app.api import quality_routes
    profile = make_profile(client)
    project, _ = plan(client, profile)
    base = f"/api/v1/me/profiles/{profile['id']}"
    monkeypatch.setattr(quality_routes, 'validate', lambda *args: {'passed': False})
    result = client.post(base + '/cases', json={'project_id': project['id'], 'label': 'Learning example', 'revision': profile['revision']})
    assert result.status_code == 201, result.text
    case = result.json(); profile = fresh(client, profile)
    assert client.post(base + f"/cases/{case['id']}/confirm", json={'revision': profile['revision']}).status_code == 409
    assert client.delete(base + f"/cases/{case['id']}").status_code == 200
    assert client.get(base + '/experience').json()['cases'] == []
    with SessionLocal() as db:
        assert db.get(PersonalBinding, project['id']) is None


def test_comparison_reveal_requires_two_reviews(client):
    profile = make_profile(client); teach(client, profile)
    project, _ = plan(client, profile)
    base = f"/api/v1/projects/{project['id']}/personal-comparisons"
    rows = client.get(base).json()
    assert len(rows) == 1 and rows[0]['revealed'] is None
    endpoint = base + '/' + rows[0]['id'] + '/review'
    scores = {key: 80 for key in ['logic', 'visual', 'completeness', 'accuracy', 'personalFit']}
    assert client.post(endpoint, json={'label': 'A', 'scores': scores}).status_code == 409
    with SessionLocal() as db:
        row = db.get(PersonalComparison, rows[0]['id']); row.status = 'rendered'; db.commit()
    assert client.post(endpoint, json={'label': 'A', 'scores': scores}).json()['revealed'] is None
    assert set(client.post(endpoint, json={'label': 'B', 'scores': scores}).json()['revealed'].values()) == {'personal', 'ordinary'}
    reset(client)
    assert client.get(base).json() == []


def test_external_markers_detect_deletions_and_text_changes(tmp_path):
    original = tmp_path / 'original.pptx'; original.write_bytes(pptx_bytes())
    annotate_export(original, 'project', [{'id': 'slide', 'content': {'title': 'Private original'}}])
    deck = Presentation(original)
    deck.slides[0].shapes.title.text = 'Changed title'
    changed = io.BytesIO(); deck.save(changed)
    aligned = compare_external(original.read_bytes(), changed.getvalue(), 'project')
    assert aligned['canSuggest'] and aligned['changes'][0]['businessTextChanged']
    shape = deck.slides[0].shapes[1]; shape._element.getparent().remove(shape._element)
    deleted = io.BytesIO(); deck.save(deleted)
    aligned = compare_external(original.read_bytes(), deleted.getvalue(), 'project')
    assert not aligned['canSuggest'] and aligned['deletedOrUnmatchedObjects'] == 1
    assert 'Changed title' not in json.dumps(aligned)


@pytest.mark.parametrize('part,data', [('ppt/vbaProject.bin', b'code'), ('ppt/embeddings/program.exe', b'code'), ('ppt/_rels/evil.rels', b'<broken')])
def test_reference_rejects_unsafe_or_broken_archives(part, data):
    buffer = io.BytesIO(pptx_bytes())
    with zipfile.ZipFile(buffer, 'a') as archive: archive.writestr(part, data)
    with pytest.raises(HTTPException) as error: validate_pptx(buffer.getvalue())
    assert error.value.status_code == 422


def test_private_accounts_isolate_projects_memories_and_revoke_sessions(client, monkeypatch):
    from app.security.public_access import _logins
    _logins.clear()
    monkeypatch.setattr(settings, 'private_accounts_mode', True)
    with TestClient(app, base_url='http://localhost', client=('127.0.0.1', 50000)) as admin:
        assert admin.post('/api/v1/auth/accounts', json={'name': 'admin', 'password': 'example-pass-123'}).status_code == 201
        assert admin.post('/api/v1/auth/login', json={'name': 'admin', 'password': 'example-pass-123'}).status_code == 200
        own = admin.post('/api/v1/projects', json={'name': 'Admin secret'}).json()
        profile = make_profile(admin)
        created = admin.post('/api/v1/auth/accounts', json={'name': 'member', 'password': 'member-pass-123'})
        assert created.status_code == 201, created.text
        with TestClient(app, base_url='http://localhost', client=('127.0.0.1', 50001)) as member:
            assert member.post('/api/v1/auth/login', json={'name': 'member', 'password': 'member-pass-123'}).status_code == 200
            assert member.get('/api/v1/projects').json() == []
            assert member.get(f"/api/v1/projects/{own['id']}").status_code == 404
            assert member.get(f"/api/v1/me/profiles/{profile['id']}/experience").status_code == 404
            assert member.post('/api/v1/projects', json={'name': 'Member project'}).status_code == 201
            assert len(admin.get('/api/v1/projects').json()) == 1
            assert admin.post('/api/v1/auth/accounts/' + created.json()['id'] + '/disable').status_code == 200
            assert member.get('/api/v1/projects').status_code == 401


def test_desktop_acceptance_is_bound_to_file_bytes(client, tmp_path, monkeypatch):
    from app.personalization import office
    from app.db.models import DeliveryVerification
    profile = make_profile(client); project, result = plan(client, profile)
    with SessionLocal() as db:
        target = Path(db.get(Project, project['id']).artifact_path) / 'exports' / 'presentation.pptx'
    target.parent.mkdir(parents=True, exist_ok=True); target.write_bytes(pptx_bytes())
    monkeypatch.setattr(office, 'render_office', lambda *args: {'status': 'unavailable', 'pages': []})
    base = f"/api/v1/projects/{project['id']}/desktop-verifications"
    response = client.post(base, json={'software': 'libreoffice'})
    assert response.status_code == 200, response.text
    verification = response.json()
    body = {'file_hash': hashlib.sha256(target.read_bytes()).hexdigest(), 'checked_fonts': True, 'checked_layout': True, 'checked_content': True}
    assert client.post(base + '/' + verification['id'] + '/accept', json=body).status_code == 409
    with SessionLocal() as db:
        row = db.get(DeliveryVerification, verification['id']); row.status = 'rendered-needs-review'; db.commit()
    target.write_bytes(b'new file')
    assert client.post(base + '/' + verification['id'] + '/accept', json=body).status_code == 409


def test_workspace_lock_reentrant_and_released_after_exception(tmp_path, monkeypatch):
    from app.personalization.process_lock import WorkspaceLock
    monkeypatch.setattr(settings, 'artifact_root', tmp_path)
    lock = WorkspaceLock()
    with pytest.raises(RuntimeError):
        with lock:
            with lock: raise RuntimeError('interrupted')
    with lock: pass

def test_native_export_requires_actual_acceptance_and_preserves_all_bullets(client, monkeypatch):
    from app.api import delivery_routes
    from app.db.models import DeliveryVerification, SlideSpecRecord
    profile = make_profile(client)
    base = f"/api/v1/me/profiles/{profile['id']}"
    response = client.post(base + '/reference-template', data={'revision': profile['revision'], 'retain_original': 'true'}, files={'file': ('sample.pptx', pptx_bytes())})
    assert response.status_code == 201
    ref = response.json(); profile = fresh(client, profile)
    client.post(base + f"/references/{ref['id']}/activate", json={'revision': profile['revision']})
    profile = fresh(client, profile); project, result = plan(client, profile)
    with SessionLocal() as db:
        slide = db.get(SlideSpecRecord, result['slides'][1]['id'])
        spec = copy.deepcopy(slide.spec); spec['content']['bullets'] = [f'Every point {index}' for index in range(8)]; slide.spec = spec; db.commit()
    endpoint = f"/api/v1/projects/{project['id']}/export/pptx"
    draft = client.get(endpoint)
    assert draft.status_code == 200, draft.text[:300] if draft.status_code != 200 else ''
    text = '\n'.join(shape.text for slide in Presentation(io.BytesIO(draft.content)).slides for shape in slide.shapes if shape.has_text_frame)
    assert all(f'Every point {index}' in text for index in range(8))
    monkeypatch.setattr(delivery_routes, 'require_final_quality', lambda *args: None)
    assert client.get(endpoint + '?stage=final').status_code == 409
    with SessionLocal() as db:
        db.add(DeliveryVerification(project_id=project['id'], file_hash=hashlib.sha256(draft.content).hexdigest(), software='powerpoint-windows', status='accepted', report={'humanVerified': True})); db.commit()
    final = client.get(endpoint + '?stage=final')
    assert final.status_code == 200 and final.content == draft.content
    with SessionLocal() as db:
        slide = db.get(SlideSpecRecord, result['slides'][1]['id']); spec = copy.deepcopy(slide.spec); spec['content']['title'] += ' changed'; slide.spec = spec; db.commit()
    assert client.get(endpoint + '?stage=final').status_code == 409


def test_comparison_renders_both_real_artifacts(client):
    profile = make_profile(client); teach(client, profile, 'font_family', 'Arial')
    project, _ = plan(client, profile, slide_count=6)
    base = f"/api/v1/projects/{project['id']}/personal-comparisons"
    row = client.get(base).json()[0]
    response = client.post(base + '/' + row['id'] + '/render')
    assert response.status_code == 200, response.text
    for label in ['A', 'B']:
        pptx = client.get(base + '/' + row['id'] + '/' + label + '/pptx')
        assert pptx.status_code == 200
        assert len(Presentation(io.BytesIO(pptx.content)).slides) == 6


def test_manual_pdf_acceptance_checks_page_content_and_hash(client):
    import pymupdf
    from app.db.models import SlideSpecRecord
    project = client.post('/api/v1/projects', json={'name': 'Manual acceptance'}).json()
    with SessionLocal() as db:
        record = db.get(Project, project['id'])
        target = Path(record.artifact_path) / 'exports' / 'presentation.pptx'
        db.add(SlideSpecRecord(id="manual-slide", project_id=record.id, position=1, spec={'id':'manual-slide','position':1,'role':'content','content':{'title':'Verified title','bullets':['Complete body']}}))
        db.commit()
    target.parent.mkdir(parents=True, exist_ok=True); target.write_bytes(pptx_bytes())
    digest = hashlib.sha256(target.read_bytes()).hexdigest()
    base = f"/api/v1/projects/{project['id']}/desktop-verifications"
    document = pymupdf.open(); page = document.new_page(); page.insert_text((70,70),'Verified title\nComplete body')
    data = document.tobytes(); document.close()
    response = client.post(base+'/manual', data={'software':'powerpoint-macos','file_hash':digest}, files={'file':('review.pdf',data)})
    assert response.status_code == 200, response.text
    result = response.json(); assert result['status'] == 'rendered-needs-review'
    assert result['report']['manualUpload'] is True
    response = client.post(base+'/'+result['id']+'/accept',json={'file_hash':digest,'checked_fonts':True,'checked_layout':True,'checked_content':True})
    assert response.status_code == 200, response.text
    assert client.post(base+'/manual', data={'software':'wps','file_hash':'outdated'}, files={'file':('review.pdf',data)}).status_code == 409


def test_additive_migrations_preserve_existing_projects(tmp_path):
    from alembic import command
    from alembic.config import Config
    from sqlalchemy import create_engine, text, inspect
    config = Config()
    config.set_main_option('script_location', str(Path('apps/api/alembic').resolve()))
    url = 'sqlite:///' + (tmp_path/'migration.db').as_posix()
    config.set_main_option('sqlalchemy.url', url)
    command.upgrade(config, '0001')
    engine = create_engine(url)
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO projects (id,name,status,artifact_path,created_at,updated_at) VALUES ('preserved','Existing','ready','test',CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)"))
    command.upgrade(config, 'head')
    assert 'personal_cases' in inspect(engine).get_table_names()
    with engine.connect() as connection: assert connection.execute(text("SELECT name FROM projects WHERE id='preserved'")).scalar() == 'Existing'
    command.downgrade(config, '0001')
    assert 'personal_cases' not in inspect(engine).get_table_names()
    with engine.connect() as connection: assert connection.execute(text("SELECT name FROM projects WHERE id='preserved'")).scalar() == 'Existing'
    engine.dispose()

